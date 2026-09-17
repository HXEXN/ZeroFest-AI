"""Shared page layout for the rasterized Korean reports in outputs/.

Korean glyphs render inconsistently across Word installations, so every report
is drawn to a PNG page and embedded as an image. This module holds the drawing
primitives the report builders share.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches

FONT_PATH = "/System/Library/Fonts/Supplemental/AppleGothic.ttf"
MONO_FONT_PATH = "/System/Library/Fonts/Supplemental/Arial.ttf"
PAGE_WIDTH, PAGE_HEIGHT = 1700, 2200
NAVY = (31, 78, 121)
GREY = (90, 90, 90)
ACCENT = (16, 150, 110)
MARGIN = 110


def fonts() -> dict[str, ImageFont.FreeTypeFont]:
    return {
        "title": ImageFont.truetype(FONT_PATH, 52),
        "heading": ImageFont.truetype(FONT_PATH, 31),
        "body": ImageFont.truetype(FONT_PATH, 23),
        "small": ImageFont.truetype(FONT_PATH, 19),
        "tiny": ImageFont.truetype(FONT_PATH, 17),
        "axis": ImageFont.truetype(MONO_FONT_PATH, 16),
    }


class Page:
    def __init__(self, font_set: dict[str, ImageFont.FreeTypeFont]) -> None:
        self.image = Image.new("RGB", (PAGE_WIDTH, PAGE_HEIGHT), "white")
        self.draw = ImageDraw.Draw(self.image)
        self.fonts = font_set
        self.y = 0

    def title(self, text: str, subtitle: str = "") -> None:
        self.draw.text((MARGIN, 95), text, font=self.fonts["title"], fill="black")
        self.draw.line((MARGIN, 178, PAGE_WIDTH - MARGIN, 178), fill=NAVY, width=4)
        self.y = 235
        if subtitle:
            self.draw.text((MARGIN, self.y), subtitle, font=self.fonts["body"], fill=GREY)
            self.y += 70

    def heading(self, text: str) -> None:
        self.y += 16
        self.draw.text((MARGIN, self.y), text, font=self.fonts["heading"], fill="black")
        self.y += 56

    def body(self, text: str, max_chars: int = 50) -> None:
        line = ""
        for word in text.split(" "):
            candidate = f"{line} {word}".strip()
            if len(candidate) > max_chars:
                self.draw.text((MARGIN, self.y), line, font=self.fonts["body"], fill="black")
                self.y += 36
                line = word
            else:
                line = candidate
        if line:
            self.draw.text((MARGIN, self.y), line, font=self.fonts["body"], fill="black")
            self.y += 36
        self.y += 12

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.draw.text((MARGIN + 18, self.y), "• " + item, font=self.fonts["body"], fill="black")
            self.y += 44
        self.y += 8

    def table(
        self, headers: list[str], rows: list[list[str]], widths: list[int],
        row_height: int = 54, highlight: int | None = None,
    ) -> None:
        for index, row in enumerate([headers] + rows):
            x = MARGIN
            if index == 0:
                fill, color = NAVY, "white"
            elif highlight is not None and index == highlight + 1:
                fill, color = (226, 244, 237), "black"
            else:
                fill = (243, 247, 250) if index % 2 == 0 else (255, 255, 255)
                color = "black"
            for value, width in zip(row, widths):
                self.draw.rectangle((x, self.y, x + width, self.y + row_height), fill=fill, outline=(205, 205, 205))
                self.draw.text((x + 14, self.y + row_height // 2 - 12), value, font=self.fonts["small"], fill=color)
                x += width
            self.y += row_height
        self.y += 26

    def note(self, text: str) -> None:
        self.draw.text((MARGIN, self.y), text, font=self.fonts["tiny"], fill=GREY)
        self.y += 32

    def paste(self, path: Path, max_size: tuple[int, int]) -> None:
        chart = Image.open(path).convert("RGB")
        chart.thumbnail(max_size)
        self.image.paste(chart, ((PAGE_WIDTH - chart.width) // 2, self.y))
        self.y += chart.height + 40

    def save(self, path: Path) -> Path:
        self.image.save(path)
        return path


def write_document(pages: list[Path], output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.3)
    section.left_margin = section.right_margin = Inches(0.3)
    width = section.page_width - section.left_margin - section.right_margin
    # Each raster page already consumes the available page body. Adding an
    # explicit page-break paragraph after it makes Word/LibreOffice place that
    # paragraph on the following page and then break again, producing a blank
    # page between every report page. Natural pagination keeps this at one
    # source image per Word page.
    for page in pages:
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(str(page), width=width)
    document.save(output)
