from __future__ import annotations

import hashlib
import math
import zipfile
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/hyeonmin/Documents/EST CHALL")
SOURCE = Path("/Users/hyeonmin/Downloads/10조_온라인해커톤_기획서초안.docx")
OUTPUT = ROOT / "ZeroFest_AI_프로젝트_기획서_최종.docx"
ASSET_DIR = ROOT / ".docx_work" / "proposal_assets"
STITCH_ZIP = Path("/Users/hyeonmin/Downloads/stitch_zerofest_ai_campus_operations_platform.zip")

INK = "1E293B"
MUTED = "5B6472"
GRID = "D9D9D9"
LABEL = "B2B2B2"
LIGHT = "F5F7FA"
EMERALD = "166534"
EMERALD_LIGHT = "E8F5EC"
ORANGE = "F97316"
ORANGE_LIGHT = "FFF1E8"
RED = "EF4444"
AMBER = "F59E0B"
BLUE = "2563EB"
SLATE = "64748B"


def set_run_font(run, size=None, bold=None, color=INK, name="Malgun Gothic"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, *, fill=None, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.1, color=INK):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if align == WD_ALIGN_PARAGRAPH.CENTER else WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)
    set_cell_border(cell)
    if fill:
        shade_cell(cell, fill)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            set_run_font(r, size=size, bold=bold, color=color)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def normalize_page_breaks(doc):
    """Convert manual break paragraphs to page-break-before on the next heading.

    This avoids an extra blank page when a preceding table naturally fills the page.
    """
    body = doc._element.body
    children = list(body)
    for idx, child in enumerate(children):
        if child.tag != qn("w:p") or not child.xpath('.//w:br[@w:type="page"]'):
            continue
        next_paragraph = None
        for candidate in children[idx + 1 :]:
            if candidate.tag == qn("w:p"):
                next_paragraph = candidate
                break
        if next_paragraph is not None:
            p_pr = next_paragraph.get_or_add_pPr()
            page_break_before = p_pr.find(qn("w:pageBreakBefore"))
            if page_break_before is None:
                page_break_before = OxmlElement("w:pageBreakBefore")
                p_pr.append(page_break_before)
        body.remove(child)


def configure_running_headers(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for header in (section.header, section.even_page_header):
            header.is_linked_to_previous = False
            root = header._element
            for child in list(root):
                root.remove(child)
            p = header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            set_run_font(p.add_run("프로젝트 기획서"), size=10, color="000000")
            p_pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "26")
            bottom.set(qn("w:space"), "5")
            bottom.set(qn("w:color"), "404040")
            pbdr.append(bottom)
            p_pr.append(pbdr)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.22

    title = doc.styles["Title"]
    title.font.name = "Malgun Gothic"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(10)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Malgun Gothic"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    subtitle.font.size = Pt(16)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if style_name == "Heading 1" else 5)
        style.paragraph_format.space_after = Pt(5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size={1: 17, 2: 13, 3: 11.5}[level], bold=True, color="000000")
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), INK)
        pbdr.append(bottom)
        p_pr.append(pbdr)
    return p


def add_body(doc, text, *, bold_lead=None, size=10.2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.22
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size)
    else:
        set_run_font(p.add_run(text), size=size)
    return p


def add_bullets(doc, items, size=9.8, color=INK, indent=0.2):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        set_run_font(p.add_run("• "), size=size, bold=True, color=EMERALD)
        set_run_font(p.add_run(item), size=size, color=color)


def add_rule(doc, color=GRID, size="8"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    p_pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    p_pr.append(pbdr)


def make_table(doc, headers, rows, widths=None, header_fill=INK, first_col_fill=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AUTO
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        if widths:
            set_cell_width(cell, widths[i])
        format_cell(cell, fill=header_fill, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size, color="FFFFFF" if header_fill != LABEL else "000000")
    for row_i, values in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            if widths:
                set_cell_width(cell, widths[i])
            fill = first_col_fill if i == 0 and first_col_fill else (LIGHT if row_i % 2 else "FFFFFF")
            format_cell(
                cell,
                fill=fill,
                bold=(i == 0 and first_col_fill is not None),
                align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 and first_col_fill else WD_ALIGN_PARAGRAPH.LEFT,
                size=font_size,
                color="000000",
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def make_key_value_table(doc, rows, label_width=1.35, value_width=4.45, font_size=9.2):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in rows:
        row = table.add_row()
        prevent_row_split(row)
        left, right = row.cells
        set_cell_width(left, label_width)
        set_cell_width(right, value_width)
        left.text = label
        right.text = value
        format_cell(left, fill=LABEL, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size, color="000000")
        format_cell(right, fill="FFFFFF", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=font_size, color="000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, image_path, caption, width=5.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.paragraph_format.space_after = Pt(5)
    set_run_font(cap.add_run(caption), size=8.5, color=MUTED)


def page_break(doc):
    doc.add_page_break()


def extract_font():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    regular = ASSET_DIR / "NanumGothic-regular.ttf"
    bold = ASSET_DIR / "NanumGothic-bold.ttf"
    if not regular.exists() or not bold.exists():
        with zipfile.ZipFile(SOURCE) as zf:
            regular.write_bytes(zf.read("word/fonts/NanumGothic-regular.ttf"))
            bold.write_bytes(zf.read("word/fonts/NanumGothic-bold.ttf"))
    return {"regular": regular, "bold": bold}


def pil_font(fonts, size, bold=False):
    return ImageFont.truetype(str(fonts["bold" if bold else "regular"]), size=size)


def canvas(width=1800, height=1000):
    img = Image.new("RGB", (width, height), "white")
    return img, ImageDraw.Draw(img)


def draw_text(draw, xy, text, fonts, size=34, fill="#1E293B", bold=False, anchor="la", spacing=10):
    draw.multiline_text(xy, text, font=pil_font(fonts, size, bold), fill=fill, anchor=anchor, spacing=spacing, align="center" if "m" in anchor else "left")


def draw_box(draw, rect, text, fonts, fill="#FFFFFF", outline="#CBD5E1", text_fill="#1E293B", size=30, bold=False, radius=18, width=3):
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = rect
    draw_text(draw, ((x1 + x2) / 2, (y1 + y2) / 2), text, fonts, size=size, fill=text_fill, bold=bold, anchor="mm", spacing=8)


def draw_arrow(draw, start, end, fill="#64748B", width=5):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 18
    spread = 0.55
    p1 = (x2 - length * math.cos(angle - spread), y2 - length * math.sin(angle - spread))
    p2 = (x2 - length * math.cos(angle + spread), y2 - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=fill)


def save_image(img, name):
    path = ASSET_DIR / name
    img.save(path, format="PNG", optimize=True)
    return path


def create_architecture(fonts):
    img, draw = canvas(1900, 1060)
    draw_text(draw, (75, 62), "ZeroFest AI 폐루프 아키텍처", fonts, size=50, bold=True)
    draw_text(draw, (75, 126), "예측을 운영 승인과 학생 수요 변화까지 연결하고, 새 판매 데이터로 다시 예측합니다.", fonts, size=27, fill="#64748B")
    sources = [
        (70, 280, 365, 410, "과거 축제 이력\n2023–2025 합성 CSV"),
        (70, 480, 365, 610, "당일 판매·재고\n30분 단위 상태"),
        (70, 680, 365, 810, "날씨·공연 일정\n외부 맥락"),
    ]
    for x1, y1, x2, y2, text in sources:
        draw_box(draw, (x1, y1, x2, y2), text, fonts, fill="#F8FAFC", size=27, bold=True)
        draw_arrow(draw, (x2, (y1 + y2) // 2), (470, 545), fill="#94A3B8")
    draw_box(draw, (475, 445, 725, 645), "Feature Pipeline\n검증·파생변수", fonts, fill="#EFF6FF", outline="#2563EB", text_fill="#1D4ED8", size=28, bold=True)
    draw_box(draw, (790, 445, 1040, 645), "수요예측 ML\nGradient Boosting", fonts, fill="#ECFDF5", outline="#16A34A", text_fill="#166534", size=28, bold=True)
    draw_box(draw, (1105, 445, 1360, 645), "폐기위험 규칙\nLOW·MEDIUM·HIGH", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#C2410C", size=27, bold=True)
    draw_box(draw, (1425, 445, 1765, 645), "LangGraph Agent\n행동 후보 생성", fonts, fill="#F5F3FF", outline="#7C3AED", text_fill="#6D28D9", size=28, bold=True)
    for start, end in [((725, 545), (790, 545)), ((1040, 545), (1105, 545)), ((1360, 545), (1425, 545))]:
        draw_arrow(draw, start, end)
    draw_box(draw, (1285, 790, 1615, 935), "운영자 승인\n미승인 시 실행 없음", fonts, fill="#FEF2F2", outline="#EF4444", text_fill="#991B1B", size=28, bold=True)
    draw_box(draw, (855, 790, 1195, 935), "승인 Action 실행\n할인·조리중단·이동", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#9A3412", size=27, bold=True)
    draw_box(draw, (425, 790, 765, 935), "학생 혜택 노출\n수요 유도", fonts, fill="#ECFDF5", outline="#10B981", text_fill="#065F46", size=28, bold=True)
    draw_arrow(draw, (1600, 645), (1450, 790), fill="#7C3AED")
    draw_arrow(draw, (1285, 862), (1195, 862), fill="#F97316")
    draw_arrow(draw, (855, 862), (765, 862), fill="#10B981")
    draw.line([(425, 862), (250, 862), (250, 810)], fill="#166534", width=5)
    draw_arrow(draw, (250, 810), (365, 745), fill="#166534")
    draw_text(draw, (75, 985), "역할 분리  ML은 수치 예측 · 규칙은 위험 판정 · Agent는 절차 관리 · 사람은 실행 결정 · LLM은 근거 설명", fonts, size=25, fill="#475569")
    return save_image(img, "architecture.png")


def create_ml_pipeline(fonts):
    img, draw = canvas(1900, 1030)
    draw_text(draw, (75, 60), "데이터 및 모델 운영 파이프라인", fonts, size=50, bold=True)
    stages = [
        ("01  INGEST", "CSV·SQLite"), ("02  VALIDATE", "Null·중복·Target"),
        ("03  FEATURE", "10개 변수"), ("04  TRAIN", "GBR"),
        ("05  EVALUATE", "시간 Hold-out"), ("06  REGISTER", "사람 검토"),
    ]
    x, y, w, gap = 65, 205, 275, 30
    for i, (name, detail) in enumerate(stages):
        fc = "#ECFDF5" if i in (3, 5) else "#F8FAFC"
        ec = "#16A34A" if i in (3, 5) else "#CBD5E1"
        draw_box(draw, (x, y, x + w, y + 155), f"{name}\n{detail}", fonts, fill=fc, outline=ec, size=25, bold=True)
        if i < 5:
            draw_arrow(draw, (x + w, y + 78), (x + w + gap, y + 78))
        x += w + gap
    draw_text(draw, (75, 435), "시간 순서 검증", fonts, size=34, bold=True)
    draw_box(draw, (75, 510, 590, 675), "2023–2024\n학습 144행", fonts, fill="#EFF6FF", outline="#2563EB", text_fill="#1D4ED8", size=34, bold=True)
    draw_arrow(draw, (590, 592), (690, 592), fill="#2563EB")
    draw_box(draw, (690, 510, 1080, 675), "2025\n검증 72행", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#C2410C", size=34, bold=True)
    draw_arrow(draw, (1080, 592), (1180, 592), fill="#166534")
    draw_box(draw, (1180, 510, 1790, 675), "실운영\n과거 이력 + 당일 신호", fonts, fill="#ECFDF5", outline="#16A34A", text_fill="#166534", size=32, bold=True)
    draw_box(draw, (75, 755, 565, 905), "MAE 1.94\n개/30분", fonts, fill="#F8FAFC", outline="#64748B", size=38, bold=True)
    draw_box(draw, (705, 755, 1195, 905), "R² 0.877\n시간 Hold-out", fonts, fill="#F8FAFC", outline="#64748B", size=38, bold=True)
    draw_box(draw, (1335, 755, 1825, 905), "216행 · 3개 가상 대학\nSample / Synthetic", fonts, fill="#FEF2F2", outline="#EF4444", text_fill="#991B1B", size=29, bold=True)
    draw_text(draw, (75, 965), "지표는 합성 데이터 내부 검증 결과이며 실제 대학 현장 성능을 의미하지 않습니다.", fonts, size=26, fill="#64748B")
    return save_image(img, "ml_pipeline.png")


def create_offline_online_workflow(fonts):
    img, draw = canvas(1900, 1110)
    draw_text(draw, (70, 55), "과거 학습과 당일 재예측을 연결하는 AI 운영 흐름", fonts, size=48, bold=True)
    draw_text(draw, (70, 118), "과거 이력은 기본 수요 패턴을 학습하고, 당일 신호는 현재 상황을 반영해 예측을 계속 갱신합니다.", fonts, size=27, fill="#64748B")

    draw_text(draw, (70, 215), "OFFLINE  모델 검증과 등록", fonts, size=31, bold=True, fill="#1D4ED8")
    offline = [
        (70, 270, 345, 405, "2023–2024\n학습 144행"),
        (410, 270, 685, 405, "검증·특징 생성\n10개 Feature"),
        (750, 270, 1025, 405, "Gradient Boosting\nHuber Loss"),
        (1090, 270, 1365, 405, "2025 Hold-out\n72행 평가"),
        (1430, 270, 1810, 405, "사람 검토·등록\n전체 이력 재학습"),
    ]
    for i, (x1, y1, x2, y2, text) in enumerate(offline):
        draw_box(draw, (x1, y1, x2, y2), text, fonts, fill="#EFF6FF", outline="#2563EB", text_fill="#1D4ED8", size=25, bold=True)
        if i < len(offline) - 1:
            draw_arrow(draw, (x2, 338), (offline[i + 1][0], 338), fill="#2563EB")

    draw_text(draw, (70, 505), "ONLINE  행사 중 추론과 사람 승인", fonts, size=31, bold=True, fill="#166534")
    online = [
        (70, 560, 350, 700, "당일 신호\n판매·재고·날씨·공연"),
        (420, 560, 700, 700, "Feature Snapshot\n시점·출처 고정"),
        (770, 560, 1050, 700, "30분·60분·종료\n판매량 예측"),
        (1120, 560, 1400, 700, "잔여비율 규칙\nLOW·MEDIUM·HIGH"),
        (1470, 560, 1810, 700, "Action 후보\n승인 전 실행 금지"),
    ]
    for i, (x1, y1, x2, y2, text) in enumerate(online):
        draw_box(draw, (x1, y1, x2, y2), text, fonts, fill="#ECFDF5", outline="#16A34A", text_fill="#166534", size=24, bold=True)
        if i < len(online) - 1:
            draw_arrow(draw, (x2, 630), (online[i + 1][0], 630), fill="#16A34A")

    draw_box(draw, (1250, 790, 1580, 920), "운영자 승인\nAction 실행", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#9A3412", size=26, bold=True)
    draw_box(draw, (820, 790, 1150, 920), "학생 혜택 노출\n수요 반응", fonts, fill="#F5F3FF", outline="#7C3AED", text_fill="#6D28D9", size=26, bold=True)
    draw_box(draw, (390, 790, 720, 920), "판매·재고 갱신\n다음 재예측", fonts, fill="#F8FAFC", outline="#64748B", text_fill="#334155", size=26, bold=True)
    draw_arrow(draw, (1640, 700), (1420, 790), fill="#F97316")
    draw_arrow(draw, (1250, 855), (1150, 855), fill="#7C3AED")
    draw_arrow(draw, (820, 855), (720, 855), fill="#64748B")
    draw.line([(390, 855), (240, 855), (240, 700)], fill="#64748B", width=5)
    draw_arrow(draw, (240, 700), (350, 650), fill="#64748B")

    draw_text(draw, (70, 1010), "설명 계층", fonts, size=28, bold=True)
    draw_text(draw, (250, 1010), "LLM 또는 Local Grounded Chat은 저장된 예측·근거·상태만 설명하며 숫자 예측과 실행 권한을 갖지 않습니다.", fonts, size=25, fill="#475569")
    draw_text(draw, (70, 1060), "회복성", fonts, size=28, bold=True)
    draw_text(draw, (250, 1060), "외부 API 장애 시 Mock Weather, Local Chat, Sequential Workflow로 같은 순서와 통제를 유지합니다.", fonts, size=25, fill="#475569")
    return save_image(img, "offline_online_workflow.png")


def create_agent_flow(fonts):
    img, draw = canvas(1700, 1160)
    draw_text(draw, (70, 58), "LangGraph Human-in-the-loop 워크플로", fonts, size=48, bold=True)
    nodes = ["1  현재 상태 로드", "2  수요예측", "3  폐기위험 계산", "4  운영조건 확인", "5  행동 후보 생성", "6  운영자 승인"]
    y = 170
    for i, text in enumerate(nodes):
        fc, ec, tc = ("#FEF2F2", "#EF4444", "#991B1B") if i == 5 else ("#F8FAFC", "#CBD5E1", "#1E293B")
        draw_box(draw, (120, y, 680, y + 105), text, fonts, fill=fc, outline=ec, text_fill=tc, size=30, bold=True)
        if i < 5:
            draw_arrow(draw, (400, y + 105), (400, y + 140))
        y += 140
    draw_box(draw, (980, 870, 1510, 1000), "미승인 · 대기\n자동 실행 없음", fonts, fill="#F8FAFC", outline="#64748B", size=29, bold=True)
    draw_arrow(draw, (680, 922), (980, 922))
    draw_text(draw, (810, 885), "NO", fonts, size=24, bold=True, fill="#64748B")
    draw_box(draw, (95, 1040, 440, 1130), "Action 실행", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#9A3412", size=29, bold=True)
    draw_arrow(draw, (310, 975), (310, 1040), fill="#F97316")
    draw_text(draw, (345, 1012), "YES", fonts, size=24, bold=True, fill="#C2410C")
    draw_box(draw, (530, 1040, 875, 1130), "결과 모니터링", fonts, fill="#ECFDF5", outline="#10B981", text_fill="#065F46", size=28, bold=True)
    draw_arrow(draw, (440, 1085), (530, 1085), fill="#10B981")
    draw_box(draw, (965, 1040, 1310, 1130), "재예측 · 종료", fonts, fill="#EFF6FF", outline="#2563EB", text_fill="#1D4ED8", size=28, bold=True)
    draw_arrow(draw, (875, 1085), (965, 1085), fill="#2563EB")
    draw_text(draw, (900, 210), "Guardrails", fonts, size=38, bold=True)
    guards = ["예측 수치는 ML만 생성", "위험도는 고정 임계값", "승인 ID 없으면 PENDING", "승인·실행 이력을 기록", "장애 시 동일 순서 fallback"]
    for i, item in enumerate(guards):
        draw_text(draw, (900, 285 + i * 88), f"• {item}", fonts, size=29, fill="#475569")
    return save_image(img, "agent_flow.png")


def create_role_ui(fonts):
    img, draw = canvas(1900, 1090)
    draw_text(draw, (75, 60), "역할 기반 화면 및 인간공학 설계", fonts, size=50, bold=True)
    panels = [
        (70, 190, 890, 570, "학생회 Control Tower", "부스 위험 비교\n예측 근거·Action Queue\nGrounded AI Chat", "#EFF6FF", "#2563EB"),
        (1010, 190, 1830, 570, "부스 운영자", "56px 원클릭 입력\n판매·재고 즉시 반영\n승인 기반 Action", "#FFF7ED", "#F97316"),
        (70, 650, 890, 1015, "학생", "할인가 우선 노출\n구역 안내·혜택\nQuiz·Stamp 확장", "#ECFDF5", "#10B981"),
        (1010, 650, 1830, 1015, "AI · ML · Data Ops", "Pipeline Control\nData·Model Registry\nAgent Audit", "#F5F3FF", "#7C3AED"),
    ]
    for x1, y1, x2, y2, title, detail, fc, ec in panels:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fc, outline=ec, width=4)
        draw_text(draw, (x1 + 35, y1 + 42), title, fonts, size=38, bold=True)
        draw_text(draw, (x1 + 35, y1 + 125), detail, fonts, size=31, fill="#475569", spacing=20)
        draw_text(draw, (x2 - 35, y2 - 38), "상태·텍스트 병기", fonts, size=23, bold=True, fill=ec, anchor="ra")
    draw_text(draw, (75, 1058), "공통 원칙  역할 격리 · 큰 터치 타깃 · 색상+텍스트 이중부호 · 실행 전 확인 · 키보드 포커스 · Reduced Motion", fonts, size=25, fill="#475569", anchor="ls")
    return save_image(img, "role_ui.png")


def create_stitch_mvp_board(fonts):
    """Arrange the four supplied Stitch screens without altering their UI content."""
    screen_specs = [
        ("zerofest_ai_1", "통합 게이트웨이"),
        ("zerofest_ai_2", "학생회 관제"),
        ("zerofest_ai_pos", "부스 운영 POS"),
        ("zerofest_ai_3", "학생 혜택·맵"),
    ]
    if not STITCH_ZIP.exists():
        raise FileNotFoundError(f"Missing supplied Stitch archive: {STITCH_ZIP}")

    screens = []
    with zipfile.ZipFile(STITCH_ZIP) as zf:
        for folder, label in screen_specs:
            member = next(name for name in zf.namelist() if name.endswith(f"/{folder}/screen.png"))
            with Image.open(BytesIO(zf.read(member))) as source:
                screen = source.convert("RGB")
            target_h = 1460
            target_w = round(screen.width * target_h / screen.height)
            screens.append((screen.resize((target_w, target_h), Image.Resampling.LANCZOS), label))

    margin, gap, label_h, footer_h = 55, 34, 92, 90
    board_w = margin * 2 + sum(screen.width for screen, _ in screens) + gap * (len(screens) - 1)
    board_h = margin + label_h + 1460 + footer_h
    board = Image.new("RGB", (board_w, board_h), "#EEF2F7")
    draw = ImageDraw.Draw(board)
    x = margin
    for screen, label in screens:
        draw.rounded_rectangle(
            (x - 8, margin + label_h - 8, x + screen.width + 8, margin + label_h + screen.height + 8),
            radius=18,
            fill="#CBD5E1",
        )
        board.paste(screen, (x, margin + label_h))
        draw_text(
            draw,
            (x + screen.width / 2, margin + 35),
            label,
            fonts,
            size=28,
            bold=True,
            fill="#0F172A",
            anchor="mm",
        )
        x += screen.width + gap
    draw_text(
        draw,
        (board_w / 2, board_h - 42),
        "첨부된 Stitch UI/UX 시안 원본 화면 · 역할별 정보 구조와 핵심 인터랙션",
        fonts,
        size=23,
        fill="#475569",
        anchor="mm",
    )
    return save_image(board, "stitch_mvp_board.png")


def create_demo_chart(fonts):
    img, draw = canvas(1600, 920)
    draw_text(draw, (70, 55), "Demo Simulation Before / After", fonts, size=49, bold=True)
    left, right, top, bottom = 190, 1490, 170, 735
    draw.line((left, bottom, right, bottom), fill="#94A3B8", width=3)
    for val in (0, 20, 40, 60, 80):
        y = bottom - int((val / 80) * (bottom - top))
        draw.line((left, y, right, y), fill="#E2E8F0", width=2)
        draw_text(draw, (left - 25, y), str(val), fonts, size=24, fill="#64748B", anchor="rm")
    bars = [(430, 70, "#EF4444", "개입 전\nHIGH"), (1010, 15, "#10B981", "개입 후\nLOW")]
    for x, val, color, label in bars:
        h = int((val / 80) * (bottom - top))
        draw.rounded_rectangle((x, bottom - h, x + 300, bottom), radius=12, fill=color)
        draw_text(draw, (x + 150, bottom - h - 42), f"{val}개", fonts, size=40, bold=True, anchor="mm")
        draw_text(draw, (x + 150, bottom + 65), label, fonts, size=30, bold=True, anchor="mm")
    draw_text(draw, (820, 350), "20% 할인 승인\n+ 학생 반응 시뮬레이션", fonts, size=29, bold=True, fill="#9A3412", anchor="mm")
    draw_arrow(draw, (930, 410), (1110, 590), fill="#F97316", width=6)
    draw_text(draw, (70, 870), "70 → 15개(55개 감소, 78.6%)는 고정 Demo Simulation 결과이며 실제 현장 효과가 아닙니다.", fonts, size=26, fill="#64748B")
    return save_image(img, "demo_chart.png")


def create_problem_solution_logic(fonts):
    img, draw = canvas(1900, 980)
    draw_text(draw, (70, 55), "문제에서 검증 가능한 결과까지의 변화 논리", fonts, size=48, bold=True)
    draw_text(draw, (70, 118), "폐기량을 사후 집계하는 대신, 행동 가능한 시점의 의사결정을 바꾸고 결과를 다시 학습합니다.", fonts, size=27, fill="#64748B")

    draw_text(draw, (75, 220), "현장 문제", fonts, size=31, bold=True, fill="#991B1B")
    problems = [
        (70, 275, 365, 415, "수요 급변\n날씨·공연·시간"),
        (420, 275, 715, 415, "부스별 정보 고립\n전체 위험 비교 불가"),
        (770, 275, 1065, 415, "늦은 잔여 인지\n행동 시간 상실"),
    ]
    for i, rect_text in enumerate(problems):
        x1, y1, x2, y2, text_value = rect_text
        draw_box(draw, (x1, y1, x2, y2), text_value, fonts, fill="#FEF2F2", outline="#EF4444", text_fill="#991B1B", size=25, bold=True)
        if i < len(problems) - 1:
            draw_arrow(draw, (x2, 345), (problems[i + 1][0], 345), fill="#EF4444")
    draw_box(draw, (1160, 275, 1815, 415), "결과\n임기응변 할인 · 과잉 재고 · 폐기", fonts, fill="#FFF7ED", outline="#F97316", text_fill="#9A3412", size=27, bold=True)
    draw_arrow(draw, (1065, 345), (1160, 345), fill="#F97316")

    draw_text(draw, (75, 525), "ZeroFest AI 개입", fonts, size=31, bold=True, fill="#166534")
    stages = [
        (70, 585, 330, 715, "1 예측\n종료 예상 잔여"),
        (385, 585, 645, 715, "2 판단\n위험·근거"),
        (700, 585, 960, 715, "3 승인\n사람이 결정"),
        (1015, 585, 1275, 715, "4 수요 전환\n학생 혜택"),
        (1330, 585, 1815, 715, "5 재예측\n실제 반응을 다음 판단으로"),
    ]
    for i, rect_text in enumerate(stages):
        x1, y1, x2, y2, text_value = rect_text
        draw_box(draw, (x1, y1, x2, y2), text_value, fonts, fill="#ECFDF5", outline="#16A34A", text_fill="#166534", size=24, bold=True)
        if i < len(stages) - 1:
            draw_arrow(draw, (x2, 650), (stages[i + 1][0], 650), fill="#16A34A")

    draw_text(draw, (75, 805), "검증 결과", fonts, size=30, bold=True)
    draw_text(draw, (75, 865), "선행지표  예측 리드타임 · 입력 완결률 · Action 승인율 · 학생 혜택 전환율", fonts, size=25, fill="#475569")
    draw_text(draw, (75, 920), "결과지표  예상 잔여 감소 · 실제 폐기 kg · 판매율 · 운영자 판단시간 · 모델 오차", fonts, size=25, fill="#475569")
    return save_image(img, "problem_solution_logic.png")


def build_document():
    source_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if source_sha != "468310fe7b26488fdb9730474e2d0094a001df7aa6d66cfc77dd4f1cf7085809":
        raise RuntimeError("Reference template changed; refusing to overwrite against an unknown source.")

    font = extract_font()
    assets = {
        "architecture": create_architecture(font),
        "ml_pipeline": create_ml_pipeline(font),
        "offline_online_workflow": create_offline_online_workflow(font),
        "agent_flow": create_agent_flow(font),
        "role_ui": create_role_ui(font),
        "stitch_mvp": create_stitch_mvp_board(font),
        "demo_chart": create_demo_chart(font),
        "problem_solution_logic": create_problem_solution_logic(font),
    }

    doc = Document(SOURCE)
    clear_body(doc)
    configure_styles(doc)
    props = doc.core_properties
    props.title = "ZeroFest AI 프로젝트 기획서"
    props.subject = "대학축제 AI 식음료 운영 플랫폼"
    props.author = "astra"
    props.keywords = "ZeroFest AI, 대학축제, 수요예측, LangGraph, Human-in-the-loop, 음식물 폐기 예방"
    props.comments = "EST AI Challengers 온라인 해커톤 제출용 완성본"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(68)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("ZeroFest AI"), size=35, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("대학축제 AI 식음료 운영 플랫폼"), size=17, bold=True, color=MUTED)
    add_rule(doc, color=INK, size="18")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(22)
    set_run_font(p.add_run("프로젝트 기획서"), size=20, bold=True, color="000000")
    make_key_value_table(
        doc,
        [
            ("주제명", "대학축제 식음료 수요·재고 예측과 승인 기반 운영 최적화"),
            ("프로젝트명", "ZeroFest AI"),
            ("팀명", "astra"),
            ("작성 기준일", "2026. 09. 10."),
        ],
        label_width=1.4,
        value_width=4.4,
        font_size=10,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    set_run_font(p.add_run("Predict  ·  Operate  ·  Engage  ·  Optimize"), size=11, bold=True, color=EMERALD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Sample / Synthetic Data 기반 MVP"), size=9.5, color=MUTED)

    # Executive summary + team
    page_break(doc)
    add_heading(doc, "기획 요약", 1)
    make_key_value_table(
        doc,
        [
            ("문제", "판매·재고·날씨·공연 정보가 분리되어 남을 재고를 행동 가능한 시점에 알기 어렵다. 핵심은 수요 변동보다 늦은 인지와 결정 지연이다."),
            ("해결", "과거 이력과 당일 신호로 종료 잔여를 갱신하고, 위험·근거·행동 후보를 운영자 승인·학생 혜택·재예측으로 연결한다."),
            ("차별성", "Winnow·Leanpath의 주방 폐기 측정 및 개선, Too Good To Go의 마감 잉여 판매와 달리 축제 진행 중의 종료 잔여를 예측하고 운영자 승인·학생 수요 전환·재예측을 하나의 캠퍼스 운영 폐루프로 연결한다."),
            ("검증", "2023–2025년 216행 합성 이력으로 시간 Hold-out 평가(MAE 1.94, R² 0.877), 자동 테스트 7건 통과. 이후 구매자·운영자·학생 인터뷰, 가격 가설, 도입의향서와 단일 대학 파일럿으로 시장성과 현장 성능을 검증한다."),
        ],
        font_size=8.9,
    )
    add_heading(doc, "1. 팀 정보", 1)
    add_heading(doc, "(1) 팀 개요", 2)
    make_key_value_table(
        doc,
        [
            ("팀명", "astra"),
            ("팀원 및 역할", "김준혁(팀장·발표), 김승은(기획), 최현민(기획·개발), 오다교(개발), 기장선(디자인)"),
            ("협업 구조", "문제 정의와 요구사항을 기획이 정리하고, 디자인이 역할별 인터페이스를 설계하며, 개발이 데이터·AI·서비스를 연결한 뒤 발표·시연팀이 심사 동선으로 검증한다."),
        ],
        font_size=9.2,
    )
    add_heading(doc, "(2) 참여 인원", 2)
    make_table(
        doc,
        ["이름", "담당 역할", "핵심 산출물"],
        [
            ["김준혁", "주제 아이디어·자료 조사·발표·시연", "문제 맥락, 3–5분 Demo 전달"],
            ["김승은", "자료 조사·기획서 작성", "문제 정의, 사회적 가치, 기획 논리"],
            ["최현민", "기획서 작성·사이트 개발", "요구사항, 전체 서비스 통합"],
            ["오다교", "사이트 개발·발표 자료 제작", "MVP 기능, 기술 시각화"],
            ["기장선", "사이트 디자인·발표 자료 제작", "역할별 UI/UX, 시연 화면"],
        ],
        widths=[0.8, 2.25, 2.75],
        font_size=8.25,
    )

    # Problem
    page_break(doc)
    add_heading(doc, "2. 문제 정의", 1)
    make_key_value_table(
        doc,
        [
            ("제안 배경 및 필요성", "팀원 김준혁은 판촉 행사 아르바이트 중 행사 종료 후 판매되지 못한 식음료가 폐기되는 현장을 보고, 이를 폐기 전에 줄일 방법을 고민했다. 대학축제 수요는 시간대·날씨·공연 종료에 따라 급변하지만 판매 둔화를 확인할 때는 조리·판촉을 바꿀 시간이 부족하다. 이 현장 경험에서 출발해 과거 이력과 당일 신호로 30분 단위 잔여를 예측하고 운영 결정을 앞당기는 ZeroFest AI를 제안한다."),
            ("타깃 대상", "핵심 사용자는 식음료 부스 운영자다. 보조 사용자는 축제 전체 위험을 조정하는 학생회 관리자, 승인된 할인과 위치 혜택을 확인하는 학생, 데이터·모델·Agent를 관리하는 AI 운영자다."),
            ("문제의 구조", "수요 급변과 부스별 정보 고립 → 품절을 피하기 위한 보수적 과잉 준비 → 늦은 잔여 인지 → 조리·이동·판촉 변경 시간 상실 → 임기응변 할인 또는 폐기라는 연쇄가 발생한다. 핵심 문제는 단순 예측 부재가 아니라 예측을 실행 가능한 결정과 학습으로 연결하는 운영 체계의 부재다."),
            ("검증 질문", "① 종료 예상 잔여를 운영자가 행동 가능한 시간 전에 제시할 수 있는가 ② 승인된 Action이 학생 화면에 즉시 반영되는가 ③ 개입 후 새 판매 데이터로 위험도가 낮아지는가 ④ 입력 부담이 혼잡한 현장에서 감당 가능한가"),
        ],
        font_size=8.85,
    )
    add_heading(doc, "유사 접근과의 차이", 2)
    make_table(
        doc,
        ["접근", "주요 기능", "한계", "ZeroFest AI의 보완"],
        [
            ["폐기 분석", "폐기량 기록·사후 리포트", "행사 중 즉시 행동 연결이 약함", "폐기 전에 예상 잔여와 Action을 제시"],
            ["POS·재고관리", "판매·재고 집계", "외부 요인과 종료 잔여 예측이 제한적", "날씨·공연·시간을 결합한 수요예측"],
            ["쿠폰·지도", "혜택 노출·방문 유도", "공급 위험과 분리됨", "고위험 재고와 학생 수요를 직접 연결"],
            ["ZeroFest AI", "예측·승인·학생 혜택·재예측", "초기 실데이터 부족", "합성 MVP 후 단일 대학 파일럿으로 검증"],
        ],
        widths=[1.0, 1.55, 1.55, 1.7],
        font_size=8.2,
    )
    add_body(doc, "정량 근거 원칙: 본 문서는 출처가 검증되지 않은 외부 통계를 성과처럼 사용하지 않는다. 현재 제시하는 수치의 범위를 합성 데이터 내부 검증과 Demo Simulation으로 제한하고, 실제 음식물 저감은 파일럿에서 kg 단위로 측정한다.", bold_lead="정량 근거 원칙:", size=9.4)

    page_break(doc)
    add_heading(doc, "문제 구조와 검증 가설", 2)
    add_heading(doc, "증상이 아닌 의사결정 실패를 해결한다", 2)
    make_table(
        doc,
        ["구조", "현장 상황", "운영 실패", "검증할 지표"],
        [
            ["수요 변동", "시간대·비·공연 종료·가격에 따라 짧은 간격으로 판매가 바뀐다.", "경험만으로 종료 수요를 판단해 과소·과잉 준비가 반복된다.", "30분 MAE, 시간대·메뉴별 오차"],
            ["정보 고립", "판매·재고·행사 맥락이 부스와 도구마다 분산된다.", "학생회가 전체 위험을 비교하지 못하고 지원 우선순위를 정하기 어렵다.", "입력 완결률, 전체 현황 갱신 지연"],
            ["인지 지연", "남은 재고를 확인할 때는 종료가 가까워져 있다.", "조리 중단·재고 이동·판촉을 실행할 시간이 사라진다.", "위험 감지 리드타임, Action 처리시간"],
            ["행동 단절", "예측·할인·학생 안내가 각각 다른 흐름으로 운영된다.", "추천이 실행되지 않거나 실행 결과가 다음 판단에 반영되지 않는다.", "승인율, 혜택 전환율, 재예측 완료율"],
            ["책임 불명확", "AI 추천의 근거와 승인 주체가 기록되지 않는다.", "잘못된 가격·조리 결정의 원인을 추적하기 어렵다.", "무승인 실행 건수, Audit 완결률"],
        ],
        widths=[0.9, 1.8, 2.0, 1.1],
        font_size=7.75,
    )
    add_heading(doc, "사용자별 해결해야 할 일", 2)
    make_table(
        doc,
        ["사용자", "상황", "완수해야 할 일", "현재 마찰"],
        [
            ["학생회", "행사 후반 위험 부스가 동시에 발생", "어디에 먼저 개입할지 비교하고 근거를 확인한다.", "현황 취합이 늦고 공통 기준이 없다."],
            ["부스 운영자", "혼잡 중 판매·조리·재고를 함께 관리", "최소 입력으로 현재 위험을 알고 실행 여부를 결정한다.", "기록 부담과 판단 피로가 크다."],
            ["학생", "축제 동선 중 구매를 결정", "가까운 혜택의 위치·가격·대기를 빠르게 확인한다.", "혜택 정보가 늦거나 분산된다."],
            ["AI 운영자", "데이터·모델·Agent 오류를 감시", "출처·성능·Fallback·승인 이력을 추적한다.", "운영 화면과 기술 상태가 섞이거나 보이지 않는다."],
        ],
        widths=[0.9, 1.55, 2.2, 1.15],
        font_size=7.8,
    )
    page_break(doc)
    add_heading(doc, "파일럿 검증 가설과 문제 경계", 2)
    make_table(
        doc,
        ["가설", "성공 기준", "실험 방법"],
        [
            ["H1 위험을 더 일찍 알 수 있다.", "운영 Action에 필요한 리드타임을 확보하고 실제 잔여 순위와 예측 위험 순위가 일치한다.", "30분 단위 예측과 실제 잔여를 부스별 비교"],
            ["H2 입력 부담을 낮출 수 있다.", "핵심 판매 입력이 10초 이내, 누락률이 운영 허용 범위 이내다.", "현장 운영자 태스크 시간·누락률 측정"],
            ["H3 승인 혜택이 수요를 전환한다.", "노출군의 판매속도·잔여 감소가 비교군보다 개선된다.", "동일 시간대 유사 부스 또는 교차 시간창 비교"],
            ["H4 다른 대학에도 적용 가능하다.", "대학·메뉴별 오차 편차가 사전 합의 범위 이내다.", "캠퍼스 단위 Hold-out과 세그먼트 오차 분석"],
        ],
        widths=[1.55, 2.45, 1.8],
        font_size=7.65,
    )
    add_heading(doc, "문제 범위", 2)
    make_key_value_table(
        doc,
        [
            ("해결 범위", "식음료 부스의 30분 단위 판매·재고 판단과 종료 전 조리·할인·이동·홍보 의사결정이다."),
            ("제외 범위", "폐기물 수거·운반·재활용 공정, 결제 정산, 개인별 이동 추적은 본 MVP의 직접 해결 범위가 아니다."),
            ("분석 단위", "대학축제 × 부스 × 메뉴 × 30분 시점이며, 학생 개인정보 대신 익명 집계 반응을 사용한다."),
            ("성공 정의", "판매율과 학생 경험을 훼손하지 않으면서 위험 감지 시간을 앞당기고 실제 폐기 kg을 줄이는 것이다."),
        ],
        font_size=8.45,
    )

    # Idea
    page_break(doc)
    add_heading(doc, "3. 아이디어 개요", 1)
    make_key_value_table(
        doc,
        [
            ("아이디어명", "ZeroFest AI"),
            ("한 줄 요약", "대학축제 식음료 부스의 수요와 종료 예상 잔여를 예측하고, 사람이 승인한 운영 조치를 학생 혜택과 재예측으로 연결하는 AI 운영 플랫폼"),
            ("핵심 컨셉", "학생회, 부스 운영자, 학생, AI 운영자가 같은 데이터 상태를 보되 각자의 과업에 필요한 화면만 사용한다. 모델은 수요를 예측하고, 규칙은 위험을 판정하며, LangGraph는 의사결정 순서를 관리한다. 최종 실행 권한은 운영자에게 남긴다."),
        ],
        font_size=9.5,
    )
    add_heading(doc, "핵심 가치 제안", 2)
    make_table(
        doc,
        ["사용자", "기존 어려움", "제공 가치", "핵심 지표"],
        [
            ["학생회", "전체 부스 위험을 한눈에 보기 어려움", "고위험 부스와 Action Queue 통합 관제", "고위험 부스 수, 승인율"],
            ["부스 운영자", "혼잡 중 기록과 판단 부담", "큰 버튼 입력, 근거가 있는 추천, 최종 승인권", "입력 소요시간, 판매율"],
            ["학생", "혜택과 위치 정보가 분산", "승인 즉시 할인가와 구역 혜택 확인", "쿠폰 확인·사용률"],
            ["AI 운영자", "데이터·모델·Agent 상태 분산", "Pipeline, Registry, Audit를 별도 관리", "품질점수, drift, 실패율"],
        ],
        widths=[0.85, 1.65, 2.0, 1.3],
        font_size=8.2,
    )
    add_heading(doc, "제품 원칙", 2)
    add_bullets(
        doc,
        [
            "폐기된 뒤 처리하는 서비스가 아니라 남기 전에 운영을 바꾸는 서비스다.",
            "AI 예측은 근거와 한계를 함께 제시하며 사람의 승인 없이 가격이나 조리를 변경하지 않는다.",
            "예측과 학생 혜택을 분리하지 않고, 실제 수요 변화가 다시 모델 입력으로 돌아오게 한다.",
            "실데이터가 없을 때는 합성 데이터임을 명시하고 실제 효과로 표현하지 않는다.",
        ],
    )

    # Solution architecture
    page_break(doc)
    add_heading(doc, "4. 문제 해결 방안", 1)
    add_heading(doc, "핵심 기능", 2)
    make_table(
        doc,
        ["기능", "입력", "AI 또는 서비스 처리", "출력·행동"],
        [
            ["1. 수요·잔여 예측", "판매·재고·시간·날씨·공연·가격", "30분·60분·종료 시점 판매량과 잔여를 예측하고 비율 기반 위험 판정", "예상 잔여, LOW·MEDIUM·HIGH, 근거 변수"],
            ["2. 승인 기반 운영", "위험도·운영조건·부스 상태", "LangGraph가 할인·조리중단·이동·프로모션 후보를 생성", "운영자가 승인한 Action만 실행·감사 기록"],
            ["3. 수요 유도·재예측", "승인된 할인과 학생 반응", "학생 화면에 혜택을 노출하고 새 판매 상태를 동일 파이프라인에 환류", "변화된 판매량, 재예측 잔여와 위험도"],
        ],
        widths=[1.05, 1.45, 2.0, 1.3],
        font_size=8.2,
    )
    add_figure(doc, assets["architecture"], "그림 1. 데이터에서 운영 승인과 학생 수요 변화까지 이어지는 ZeroFest AI 폐루프", width=5.78)
    add_body(doc, "차별점은 예측 모델 하나가 아니라 예측 결과가 현장 행동으로 전환되고, 행동 결과가 다시 데이터가 되는 전체 운영 구조에 있다.", size=9.4)

    page_break(doc)
    add_heading(doc, "해결 전략과 변화 논리", 2)
    add_figure(
        doc,
        assets["problem_solution_logic"],
        "그림 2. 현장 의사결정 실패를 예측·승인·수요 전환·재예측으로 바꾸는 변화 논리",
        width=5.78,
    )
    add_heading(doc, "폐기 예방 폐루프", 2)
    make_table(
        doc,
        ["단계", "입력과 판단", "책임 주체", "출력과 다음 단계", "안전장치"],
        [
            ["1 감지", "판매·재고·시간·날씨·공연을 동일 시점으로 결합", "데이터 Pipeline", "Feature Snapshot", "출처·시각·data_class 기록"],
            ["2 예측", "30분 판매와 종료 예상 잔여 산출", "Machine Learning", "예측량·잔여·근거", "모델 출처·Fallback 표시"],
            ["3 판단", "잔여비율과 운영조건으로 위험·후보 결정", "결정 규칙·Agent", "LOW·MEDIUM·HIGH와 Action 후보", "임계값·정책을 코드로 고정"],
            ["4 승인", "근거·영향·대상 부스를 확인하고 실행 여부 선택", "부스 운영자", "APPROVED·EXECUTED Action", "무승인·타 부스·중복 실행 차단"],
            ["5 수요 전환", "승인된 가격·위치·혜택만 학생에게 노출", "Student Service", "Promotion·방문·판매 변화", "원가·할인가·유효시간 병기"],
            ["6 학습", "새 판매·재고 상태로 같은 Workflow 재실행", "Prediction·AI Ops", "Before/After·Audit·후속 데이터", "성과와 Simulation을 분리 기록"],
        ],
        widths=[0.75, 1.65, 1.1, 1.55, 0.75],
        font_size=7.3,
    )
    add_heading(doc, "왜 이 해결 방식이 필요한가", 2)
    add_bullets(
        doc,
        [
            "예측만으로는 현장이 바뀌지 않으므로 책임자·행동·시점을 함께 제시한다.",
            "가격·조리·재고 이동은 사람의 승인과 취소 경로를 거쳐야 한다.",
            "할인은 공급 위험과 연결해 불필요한 매출 손실과 과도한 판촉을 줄인다.",
            "개입 후 실제 반응을 재예측해 효과와 실패를 다음 판단 데이터로 남긴다.",
        ],
    )

    page_break(doc)
    add_heading(doc, "차별성과 확장 경쟁력", 2)
    add_heading(doc, "대체 수단 대비 차별성", 2)
    make_table(
        doc,
        ["비교 항목", "POS·재고", "쿠폰·지도", "예측 Dashboard", "ZeroFest AI"],
        [
            ["핵심 목적", "거래·재고 기록", "방문·구매 유도", "미래 수치 제공", "폐기 전 운영 의사결정 변화"],
            ["과거+당일 맥락", "일부 판매 중심", "제한적", "가능", "판매·재고·날씨·공연·할인 결합"],
            ["종료 잔여·위험", "수동 확인", "없음", "예측 가능", "예측+결정론적 위험 임계값"],
            ["실행 책임", "사용자 별도 판단", "즉시 노출", "대개 범위 밖", "근거 확인 후 운영자 승인"],
            ["학생 수요 연결", "없음", "핵심 기능", "대개 없음", "승인된 고위험 재고만 혜택으로 연결"],
            ["개입 후 학습", "판매 기록", "사용량 집계", "재학습 가능", "Action·노출·판매·재예측 인과 로그"],
            ["AI 운영·감사", "제한적", "제한적", "모델 지표", "Data·Model·Prediction·Agent Audit"],
            ["장애 회복성", "로컬 운영 가능", "서비스 의존", "모델 의존", "Weather·LLM·LangGraph 단계별 Fallback"],
        ],
        widths=[1.15, 1.05, 1.05, 1.2, 1.35],
        font_size=7.25,
    )
    add_heading(doc, "차별성의 네 계층", 2)
    make_table(
        doc,
        ["계층", "핵심 내용", "모방 난이도를 높이는 축적 자산", "검증 방법"],
        [
            ["운영", "학생회·부스·학생·AI Ops의 역할별 의사결정 흐름", "대학축제 운영 절차와 승인 규칙", "역할별 태스크 성공률"],
            ["데이터", "예측 입력뿐 아니라 Action·승인·노출·반응을 연결", "캠퍼스·메뉴·시간대별 개입 결과 데이터", "추적성·결측률·세그먼트 커버리지"],
            ["AI", "ML 수치 예측, 규칙 위험 판정, Agent 절차, LLM 설명 분리", "Feature 계약·정책 버전·Audit 이력", "오차·무승인 실행·설명 근거율"],
            ["제품", "예측에서 학생 행동과 재예측까지 닫힌 폐루프", "반복 가능한 Demo와 POS·Weather Adapter", "End-to-end 완료율·복구 시간"],
        ],
        widths=[0.8, 2.0, 2.0, 1.0],
        font_size=7.55,
    )
    add_body(
        doc,
        "방어 가능성은 알고리즘 독점에서 주장하지 않는다. 실제 경쟁력은 여러 축제에서 축적되는 ‘상태→추천→사람 결정→학생 반응→결과’ "
        "데이터와 대학별 운영 규칙, 현장 통합 경험에서 형성된다. 현재 MVP는 이 구조를 검증하는 단계다.",
        bold_lead="방어 가능성:",
        size=8.65,
    )

    page_break(doc)
    add_heading(doc, "유사 서비스 분석과 전략적 포지셔닝", 2)
    add_body(
        doc,
        "결론: 세 서비스는 음식물 쓰레기 저감이라는 같은 목표를 갖지만 문제를 포착하는 시점과 행동 주체가 다르다. "
        "Winnow와 Leanpath는 상업용 주방의 폐기 데이터를 측정·분석해 반복 낭비를 줄이고, Too Good To Go는 이미 발생한 판매 가능 잉여를 소비자에게 연결한다. "
        "ZeroFest AI는 이들을 대체하기보다 대학축제의 판매 진행 중에 종료 잔여 위험을 예측하고, 부스 운영자의 승인과 학생 수요 전환을 연결하는 미충족 운영 구간을 담당한다.",
        bold_lead="결론:",
        size=9.35,
    )
    make_table(
        doc,
        ["서비스", "주요 고객과 범위", "핵심 방식", "상대적 범위와 ZeroFest의 포지션"],
        [
            ["Winnow", "호텔·뷔페·크루즈 등 B2B 상업용 주방", "AI 카메라와 연결 저울로 버려진 음식의 종류·무게·비용·사유를 자동 기록하고 개선 인사이트 제공", "강점은 자동 폐기 측정과 주방 분석이다. ZeroFest는 별도 전용 폐기 인식 장비 없이 행사 맥락과 판매 흐름으로 폐기 전 운영 Action을 지원한다."],
            ["Leanpath", "급식·대학 식당·호텔 등 B2B 식음 운영 조직", "전용 트래킹 도구와 분석 플랫폼으로 조리 단계의 폐기 원인과 운영 패턴을 파악하고 예방 행동 지원", "강점은 엔터프라이즈 폐기 관리와 행동 개선이다. ZeroFest는 단기 행사·다수 부스의 실시간 위험 우선순위와 승인 흐름에 집중한다."],
            ["Too Good To Go", "베이커리·마트·식당과 일반 소비자를 잇는 B2C 플랫폼", "Surprise Bag을 앱에서 예약·결제하고 정해진 시간에 수령하도록 해 판매 가능 잉여를 소비자에게 연결", "강점은 잉여 판매의 대규모 소비자 연결이다. ZeroFest는 잉여가 확정되기 전 예측을 시작하고 승인된 고위험 재고만 캠퍼스 학생 혜택으로 전환한다."],
            ["ZeroFest AI", "학생회·부스 운영자·학생·AI 운영자를 잇는 캠퍼스 B2B2C", "과거 이력과 당일 신호로 종료 잔여를 예측하고 Action 후보·사람 승인·혜택 노출·재예측·감사를 연결", "MVP 단계로 현장 실증이 필요하다. 차별점은 대학축제 특화 맥락, 폐기 전 리드타임, 다부스 관제와 개입 결과 데이터의 폐루프다."],
        ],
        widths=[0.9, 1.45, 1.85, 1.6],
        font_size=7.15,
    )
    add_heading(doc, "기능 범위 비교", 2)
    make_table(
        doc,
        ["판단 기준", "Winnow", "Leanpath", "Too Good To Go", "ZeroFest AI"],
        [
            ["주요 개입 시점", "폐기 발생 시점의 자동 기록과 이후 개선", "조리·운영 폐기 측정과 지속 개선", "마감 잉여 확정 이후 판매", "행사 진행 중 종료 잔여 위험 예측"],
            ["핵심 데이터", "폐기 이미지·무게·비용·사유", "폐기 무게·유형·운영 패턴", "판매자 등록 잉여·예약·결제", "판매·재고·날씨·공연·가격·Action"],
            ["현장 실행 연결", "주방 리포트와 개선 인사이트", "처방형 가이드와 행동 관리", "소비자 예약·수령", "근거 기반 Action 후보와 운영자 승인"],
            ["수요 전환과 환류", "소비자 수요 연결은 핵심 범위 아님", "소비자 수요 연결은 핵심 범위 아님", "소비자 수요 연결이 핵심", "학생 노출 후 판매 변화·재예측·Audit"],
        ],
        widths=[1.15, 1.1, 1.1, 1.15, 1.3],
        font_size=6.95,
    )
    add_body(
        doc,
        "전략적 해석: ZeroFest는 폐기 실측값을 학습 Label로 받아 예측을 개선하고, 마감 후 남은 판매 가능 잉여는 외부 재분배 채널로 넘길 수 있다. "
        "핵심 제품 경계는 ‘축제 중 폐기 위험을 먼저 발견하고 승인 가능한 행동으로 바꾸는 운영 계층’이다.",
        bold_lead="전략적 해석:",
        size=8.1,
    )
    add_body(
        doc,
        "공식 공개 정보 기준 2026.09.10. Winnow: winnowsolutions.com / Leanpath: leanpath.com / Too Good To Go: toogoodtogo.com. 공개 기능 범위를 기준으로 비교했으며 가격·성과 수치는 사용하지 않았다.",
        bold_lead="자료 기준:",
        size=6.8,
    )

    # ML pipeline
    page_break(doc)
    add_heading(doc, "과거 데이터 기반 예측과 당일 보정", 2)
    add_body(doc, "결론: 전년도 데이터를 포함한 과거 축제 이력으로 기본 패턴을 학습하는 것이 맞다. 다만 전년도 데이터만으로 예측하지 않는다. 동일 축제 당일의 최근 30분 판매, 직전 30분 판매, 현재 재고, 종료까지 남은 시간, 강수확률, 공연 종료 여부, 승인 할인율을 함께 넣어 급변하는 현장을 보정한다.", bold_lead="결론:", size=10.2)
    add_figure(doc, assets["ml_pipeline"], "그림 3. 과거 이력 학습, 시간 순서 검증, 당일 신호 보정으로 이어지는 ML 파이프라인", width=5.78)
    make_key_value_table(
        doc,
        [
            ("학습 데이터", "2023–2025년 3개 가상 대학의 216개 합성 과거 축제 레코드. 실제 전국 대학 데이터가 아님."),
            ("검증 방식", "2023–2024년 144행 학습, 2025년 72행 시간 Hold-out 평가로 미래 정보 누출 방지."),
            ("모델", "GradientBoostingRegressor, 120 estimators, depth 3, learning rate 0.045, Huber loss."),
            ("입력 변수", "최근·이전 30분 판매, 시간, 종료까지 분, 강수확률, 공연 종료, 할인율, 메뉴 유형, 가격, 캠퍼스 규모."),
            ("평가 결과", "MAE 1.942개/30분, R² 0.877. 합성 데이터 내부 성능이며 실제 현장 성능으로 일반화하지 않음."),
            ("위험 규칙", "예상 잔여 ÷ 현재 재고: 10% 미만 LOW, 10–30% MEDIUM, 30% 이상 HIGH."),
        ],
        font_size=8.65,
    )

    # Agent flow
    page_break(doc)
    add_heading(doc, "AI Agent 실행 구조", 2)
    add_figure(doc, assets["agent_flow"], "그림 4. 운영자 승인을 분기점으로 두는 LangGraph Human-in-the-loop 워크플로", width=5.55)
    make_table(
        doc,
        ["계층", "책임", "오작동 방지"],
        [
            ["ML Prediction", "향후 판매량과 종료 예상 잔여", "LLM이 숫자를 임의 생성하지 않음"],
            ["Waste Risk", "잔여 비율로 위험 단계 계산", "환경설정 임계값의 결정론적 규칙"],
            ["LangGraph", "상태 로드부터 재예측까지 순서·분기", "승인 ID 없으면 대기 종료"],
            ["Business Rules", "행동 후보 생성", "관측된 조건이 있는 후보만 제시"],
            ["Human Approval", "할인·중단·이동의 최종 실행 결정", "PENDING → APPROVED → EXECUTED 기록"],
            ["LLM Explanation", "DB 스냅샷을 자연어로 설명", "API 장애 시 같은 근거의 로컬 설명"],
        ],
        widths=[1.2, 2.35, 2.25],
        font_size=8.35,
    )

    # Effects
    page_break(doc)
    add_heading(doc, "5. 기대효과 및 사회적 가치", 1)
    add_figure(doc, assets["demo_chart"], "그림 5. 닭꼬치 부스 고정 시연 시나리오의 예상 잔여 변화", width=5.45)
    add_body(doc, "시연 해석: 재고 180개, 종료 예상 잔여 70개, HIGH 상태에서 운영자가 20% 마감 할인을 승인하고 학생 반응을 시뮬레이션하면 예상 잔여가 15개, LOW로 재계산된다. 55개 감소와 78.6% 변화는 기능 연결을 보여주는 고정 Demo Simulation이며, 실증 성과가 아니다.", bold_lead="시연 해석:", size=9.3)
    add_heading(doc, "파일럿 성과 지표", 2)
    make_table(
        doc,
        ["가치 영역", "핵심 KPI", "측정 방법", "MVP 기준"],
        [
            ["환경", "종료 잔여·실폐기 kg", "부스별 종료 재고와 폐기 계량", "예측·기록 구조 구현"],
            ["경제", "판매율·폐기원가·추가 매출", "POS 판매와 원가 입력 비교", "할인 후 가격·판매 상태 연결"],
            ["운영", "고위험 감지 선행시간·Action 승인율", "예측·승인 타임스탬프", "AgentAction 감사 로그 구현"],
            ["학생", "혜택 확인·사용률", "익명 쿠폰 이벤트 집계", "승인 즉시 학생 화면 노출"],
            ["AI 품질", "MAE·R²·drift·오류율", "시간 Hold-out과 캠퍼스별 모니터링", "합성 기준 MAE 1.94, R² 0.877"],
        ],
        widths=[0.9, 1.5, 2.0, 1.4],
        font_size=8.15,
    )
    add_heading(doc, "사회적 가치", 2)
    add_bullets(
        doc,
        [
            "환경: 음식과 포장재가 폐기물로 전환되기 전 수요와 공급을 조정하는 예방 중심 접근이다.",
            "경제: 운영자의 품절 회피와 잔여 손실 사이의 균형을 데이터로 지원한다.",
            "교육: 학생이 할인·Quiz·Stamp를 통해 자원순환 행동에 참여하도록 설계할 수 있다.",
            "확장: 대학 단위로 축적한 익명 운영 데이터를 축제·지역 행사 단위 모델로 확장할 수 있다.",
        ],
        size=9.2,
    )

    # Feasibility
    page_break(doc)
    add_heading(doc, "6. 실현 가능성", 1)
    make_key_value_table(
        doc,
        [
            ("필요 기술 및 리소스", "Streamlit 역할별 UI, Python·Pandas 전처리, scikit-learn 수요예측, SQLite 단일 상태 저장소, LangGraph 실행 흐름, 선택형 Open-Meteo·OpenAI Adapter."),
            ("현재 구현 상태", "게이트웨이, 학생회 관제, 부스 운영자, 학생, Before/After, AI·ML·Data Ops 화면이 연결되어 있다. 데이터 품질·학습·Registry·Agent Audit와 Demo Reset을 포함하며 자동 테스트 7건이 통과했다."),
            ("데모 회복성", "API 키 없이 핵심 흐름이 작동한다. Weather API 실패 시 버전 관리 Mock, LLM 실패 시 DB 기반 Local Grounded Chat, LangGraph 의존성 실패 시 동일 노드 순서 fallback을 사용한다."),
        ],
        font_size=9.05,
    )
    add_heading(doc, "제약과 대응", 2)
    make_table(
        doc,
        ["리스크", "영향", "대응", "검증 기준"],
        [
            ["초기 실데이터 부족", "캠퍼스 일반화 불확실", "합성 표기, 단일 대학 파일럿, 행사 후 재학습", "캠퍼스·메뉴별 MAE"],
            ["현장 입력 부담", "누락·지연 데이터", "+1/+5/+10, 재고 ±10, POS Adapter", "입력시간·누락률"],
            ["날씨·공연 급변", "예측 drift", "30분 rolling 재예측, 외부 API fallback", "drift·오류율"],
            ["AI 추천 과신", "부적절한 할인·중단", "근거 표시, 운영자 승인, 감사 로그", "무승인 실행 0건"],
            ["할인 남용", "브랜드·수익 훼손", "최대 할인율·시간창·재고 비율 규칙", "원가 이하 추천 0건"],
            ["개인정보 위험", "신뢰 저하", "정밀 GPS·불필요 개인정보 미수집", "PII 필드 0개"],
        ],
        widths=[1.1, 1.2, 2.2, 1.3],
        font_size=7.9,
    )
    add_heading(doc, "단계별 실행 계획", 2)
    make_table(
        doc,
        ["단계", "범위", "완료 조건"],
        [
            ["MVP", "합성 데이터, 3역할 운영, Agent 승인, AI Ops", "End-to-end Demo와 자동 테스트 통과"],
            ["단일 대학 파일럿", "2–3일 축제, 참여 부스, 실제 판매·잔여 수집", "현장 MAE, 폐기 kg, 승인율 산출"],
            ["다대학 확장", "Cross-campus 학습, POS·쿠폰 연동", "대학별 편차와 모델 drift 관리"],
        ],
        widths=[1.15, 2.45, 2.2],
        font_size=8.4,
    )

    page_break(doc)
    add_heading(doc, "시장검증 계획", 2)
    add_body(
        doc,
        "목표: 기술이 작동하는지와 시장이 채택하는지를 분리해 검증한다. 경제적 구매자는 학생회·대학 행사 담당 부서·축제 대행사로 가정하고, "
        "부스 운영자는 핵심 사용자, 학생은 혜택 반응 사용자로 구분한다. 아래 표의 수치는 달성 성과가 아니라 인터뷰와 파일럿에서 확인할 사전 기준이다.",
        bold_lead="목표:",
        size=9.35,
    )
    add_heading(doc, "고객과 시장 가설", 2)
    make_table(
        doc,
        ["대상", "핵심 과업과 문제 가설", "가치 가설", "검증할 증거"],
        [
            ["경제적 구매자\n학생회·대학·대행사", "부스별 폐기 위험을 조기에 파악하고 전체 운영·ESG 성과를 관리해야 한다.", "통합 관제와 사후 성과 보고가 운영 위험·취합 시간을 줄이면 행사 단위 예산을 배정한다.", "현재 관리 방식·손실 사례·예산 주체·구매 절차·도입의향서"],
            ["핵심 사용자\n부스 운영자", "혼잡 중 최소 입력으로 남은 재고와 실행 시점을 판단해야 한다.", "10초 이내 입력과 근거가 있는 Action이 판단 부담을 낮추면 행사 내내 사용한다.", "현재 판단 기준·입력 소요시간·추천 승인율·반복 사용 의향"],
            ["반응 사용자\n학생", "축제 동선에서 가까운 혜택·가격·수령 가능성을 빠르게 확인해야 한다.", "승인된 고위험 재고 혜택이 명확하면 방문과 구매로 이어진다.", "혜택 확인률·방문 의향·실제 사용률·가격 정보 이해도"],
            ["확장 파트너\nPOS·급식·행사 운영사", "판매·재고·성과 데이터를 중복 입력하지 않고 연계해야 한다.", "표준 Adapter와 결과 리포트가 있으면 다부스·다대학 확장이 가능하다.", "필수 연동 필드·API 가능 범위·운영 책임·파트너십 의향"],
        ],
        widths=[1.15, 1.7, 1.8, 1.15],
        font_size=7.55,
    )
    add_heading(doc, "검증할 핵심 시장 가설", 2)
    make_table(
        doc,
        ["가설", "검증 질문", "최소 증거"],
        [
            ["MV-01 문제 빈도", "폐기와 늦은 대응이 다음 행사에서도 반복될 중요한 문제인가?", "구매자·운영자의 70% 이상이 실제 사례와 현재 대응을 설명"],
            ["MV-02 제품 채택", "운영 흐름을 방해하지 않고 핵심 과업을 완료할 수 있는가?", "역할별 핵심 태스크 성공률 80% 이상, 판매 입력 중앙값 10초 이내"],
            ["MV-03 지불의향", "행사 단위 비용을 낼 주체와 수용 가능한 가격 범위가 존재하는가?", "구매자 3곳 이상 가격 인터뷰, 1곳 이상 유료 범위 수용"],
            ["MV-04 도입 전환", "데모 관심이 실제 데이터 제공과 파일럿 참여로 이어지는가?", "파일럿 후보 2곳 이상, 데이터 협의 1곳, 도입의향서 1건 이상"],
        ],
        widths=[1.2, 2.8, 1.8],
        font_size=7.65,
    )
    add_body(
        doc,
        "가격 검증은 무료 의향만 묻지 않는다. 행사당 50만·100만·150만원의 가설 가격 카드를 무작위 순서로 제시하고, 각 가격에서 구매·보류·거절 이유와 "
        "예산 출처를 기록한다. 이 금액은 확정 판매가가 아니며 인터뷰 후 기능 범위와 비용 구조에 맞춰 조정한다.",
        bold_lead="가격 실험:",
        size=8.25,
    )

    page_break(doc)
    add_heading(doc, "시장검증 실행 로드맵과 통과 기준", 2)
    add_heading(doc, "사전 검증부터 유료 전환까지", 2)
    make_table(
        doc,
        ["단계", "기간과 표본 목표", "실행 방법", "산출물"],
        [
            ["1 문제 탐색", "1주\n구매자 8명·운영자 12명·학생 30명", "최근 행사 기준 인터뷰로 폐기 상황·의사결정 지연·현재 도구·예산 주체를 확인한다.", "문제 사례 목록, 고객 여정, 구매위원회 지도"],
            ["2 솔루션 검증", "1주\n구매자 5명·운영자 8명·학생 10명", "Stitch 기반 MVP로 관제→승인→학생 혜택→재예측 과업을 수행하게 하고 행동·시간·오류를 관찰한다.", "태스크 성공률, 입력 시간, 이탈 지점, 개선 Backlog"],
            ["3 가격·도입 검증", "1주\n구매 조직 3곳 이상", "문제 요약과 Demo 후 가설 가격 카드를 제시하고 예산·계약·데이터 제공 조건을 확인한다.", "가격 반응표, 파일럿 제안서, 반대 사유"],
            ["4 현장 파일럿", "축제 2–3일\n1개 대학·5–10개 부스", "일부 부스에 적용하고 비교 가능한 부스·시간창을 두어 판매·잔여·폐기 kg·승인·학생 반응을 수집한다.", "현장 KPI, 인터뷰 기록, 도입의향서 또는 개선 합의"],
            ["5 전환 검증", "행사 후 1주\n구매 의사결정자", "성과 리포트와 다음 행사 제안을 제출해 유료 전환·재사용·추천 의향을 확인한다.", "유료 제안 결과, 재사용 의향, 다음 검증 결정"],
        ],
        widths=[1.0, 1.25, 2.45, 1.1],
        font_size=7.35,
    )
    add_heading(doc, "Go Pivot Stop 의사결정 게이트", 2)
    make_table(
        doc,
        ["Gate", "Go 기준", "미달 시 결정"],
        [
            ["Problem Gate", "실제 손실 사례 5건 이상, 문제 확인 응답 70% 이상", "폐기보다 더 큰 운영 문제와 고객군을 다시 정의"],
            ["Adoption Gate", "핵심 태스크 성공률 80% 이상, 입력 10초 이내, 반복 사용 의향 70% 이상", "기능 추가보다 입력·승인 동선을 단순화"],
            ["Market Gate", "파일럿 후보 2곳, 도입의향서 1건, 가설 가격 수용 조직 1곳 이상", "구매자를 학생회·대학·대행사 중 반응이 높은 집단으로 좁힘"],
            ["Impact Gate", "기준선 대비 폐기 kg와 종료 잔여가 개선되고 매출·학생 경험이 허용 범위 내 유지", "할인 정책·위험 임계값·대상 메뉴를 조정한 뒤 재실험"],
        ],
        widths=[1.15, 2.75, 1.9],
        font_size=7.55,
    )
    add_body(
        doc,
        "증거 관리: 인터뷰 발언, 태스크 로그, 가격 반응, 도입의향서, 현장 KPI를 가설 ID와 연결한다. 한 번의 파일럿 결과를 시장 전체 성과로 일반화하지 않고, "
        "반증된 가설과 변경 이력까지 기록해 다음 대학에서 반복 검증한다.",
        bold_lead="증거 관리:",
        size=8.45,
    )

    # MVP Design
    page_break(doc)
    add_heading(doc, "7. MVP 설계서", 1)
    add_figure(doc, assets["role_ui"], "그림 6. 역할별 정보 밀도와 주요 조작을 분리한 화면 설계", width=5.78)
    make_table(
        doc,
        ["화면", "핵심 과업", "주요 구성", "인간공학 기준"],
        [
            ["통합 게이트웨이", "역할 선택·데모 상태 확인", "Admin·Operator·Student·AI Ops 진입", "선택지 축소, 데이터 범위 선표시"],
            ["학생회 관제", "위험 부스 비교·근거 확인", "KPI, 위험표, Action Queue, AI Chat", "비교 중심 고밀도, 상태 텍스트 병기"],
            ["부스 운영자", "판매 입력·Action 승인", "큰 수량 버튼, 현재 재고, 추천", "최소 56px 타깃, 위험 행동 전 확인"],
            ["학생", "할인·위치·보상 확인", "할인가, Mock Map, Quiz, Stamp", "혜택 우선 저밀도, 원가·할인가 동시 표시"],
            ["AI Ops", "데이터·모델·Agent 운영", "Pipeline Control, Registry, Audit", "운영 UI와 분리, 기술 경고 상시 노출"],
        ],
        widths=[1.05, 1.45, 1.85, 1.45],
        font_size=8.0,
    )
    add_heading(doc, "MVP 범위", 2)
    make_table(
        doc,
        ["우선순위", "포함 기능", "제외 또는 다음 단계"],
        [
            ["Must", "판매·재고 입력, 수요·잔여 예측, 위험도, 승인 할인, 학생 노출, 재예측, 감사 로그", "—"],
            ["Should", "재고 이동, 조리 중단, 날씨 실 API, Grounded LLM 설명", "현장 정책 확인 후 활성화"],
            ["Later", "실시간 혼잡도, 정밀 경로, 쿠폰 정산, 다대학 모델", "개인정보·정산·POS 계약 필요"],
        ],
        widths=[1.0, 3.0, 1.8],
        font_size=8.3,
    )

    page_break(doc)
    add_heading(doc, "첨부 UI/UX 시안 기반 MVP 화면", 2)
    add_body(
        doc,
        "사용자가 제공한 Stitch 디자인을 실제 MVP 화면 기준으로 반영했다. 통합 진입, 학생회 관제, "
        "부스 POS, 학생 혜택 화면이 동일한 운영 상태를 역할에 맞는 정보 밀도로 보여준다.",
        size=9.3,
    )
    add_figure(
        doc,
        assets["stitch_mvp"],
        "그림 7. 제공된 Stitch UI/UX 시안을 적용한 ZeroFest AI 역할별 MVP 화면",
        width=5.78,
    )

    # Requirements 8-1
    page_break(doc)
    add_heading(doc, "8. 요구사항정의서", 1)
    add_heading(doc, "(8-1) 핵심 기능 상세", 2)
    make_table(
        doc,
        ["기능명", "사용자 행동", "AI 또는 서비스가 하는 일"],
        [
            ["수요·잔여 예측", "관리자가 부스를 선택하거나 전체 분석을 실행한다.", "판매·재고·환경 변수를 읽어 30분·60분·종료 판매량과 잔여, 위험도, 근거를 저장한다."],
            ["현장 판매 입력", "운영자가 +1/+5/+10 또는 재고 ±10을 누른다.", "판매와 재고를 한 트랜잭션으로 갱신하고 다음 분석의 현재 상태로 제공한다."],
            ["Action 승인", "운영자가 할인·조리중단·이동 후보를 확인하고 승인한다.", "승인 ID와 상태를 검증한 Action만 실행하며 PENDING·APPROVED·EXECUTED를 기록한다."],
            ["학생 혜택", "학생이 주변 혜택과 할인가를 확인한다.", "승인된 Promotion만 노출하고 원가·할인가·구역 정보를 같은 상태에서 읽는다."],
            ["재예측", "시뮬레이션 또는 실제 판매 변화 뒤 재분석한다.", "새 상태와 승인 할인율을 반영해 잔여와 위험을 갱신하고 Before/After를 비교한다."],
            ["AI Ops", "AI 운영자가 데이터 검증·학습·감사 화면을 확인한다.", "Ingest→Validate→Feature→Train→Evaluate→Register 단계와 모델·Agent 이력을 관리한다."],
        ],
        widths=[1.2, 2.05, 2.55],
        font_size=8.1,
    )
    add_heading(doc, "수용 기준", 2)
    make_table(
        doc,
        ["ID", "요구사항", "검증 방법"],
        [
            ["AC-01", "승인 전 학생 화면에 할인 가격이 나타나지 않는다.", "PENDING 상태와 Student UI 확인"],
            ["AC-02", "20% 할인 승인 후 6,000원은 4,800원으로 표시된다.", "Operator 승인 후 Student UI 확인"],
            ["AC-03", "무효·타 부스·이미 처리된 승인 ID는 실행하지 않는다.", "자동 테스트와 Audit 로그 확인"],
            ["AC-04", "학생 반응 후 예상 잔여와 위험도가 재계산된다.", "70 HIGH → 15 LOW Demo Simulation"],
            ["AC-05", "데이터 출처와 모델 출처가 화면에 표시된다.", "Admin·AI Ops 경고·Model Registry 확인"],
            ["AC-06", "API 장애에도 핵심 Demo가 완료된다.", "Mock Weather·Local Chat·Sequential fallback"],
        ],
        widths=[0.75, 3.3, 1.75],
        font_size=8.2,
    )

    page_break(doc)
    add_heading(doc, "기능 요구사항 상세 1", 2)
    add_body(
        doc,
        "요구사항은 FR 기능, NFR 비기능, DR 데이터, AIR AI 통제로 구분한다. Must는 시연 폐루프에 반드시 필요하고, "
        "Should는 파일럿 품질을 높이며, Could는 다대학 확장 단계에서 구현한다. 각 항목은 입력 조건, 처리 결과, "
        "수용 기준을 함께 정의해 개발과 테스트의 기준선을 일치시킨다.",
        size=9.0,
    )
    make_table(
        doc,
        ["ID·우선순위", "요구사항", "선행 조건과 처리 결과", "수용 기준"],
        [
            ["FR-001\nMust", "역할별 진입", "사용자가 Gateway에서 학생회·운영자·학생·AI Ops를 선택하면 독립 화면으로 이동한다.", "모든 역할 링크가 동작하고 다른 역할의 조작 기능이 섞이지 않는다."],
            ["FR-002\nShould", "부스 등록", "부스명·구역·메뉴·가격·초기재고를 입력하면 Booth와 Menu가 생성되고 초기 재고가 기록된다.", "필수값 누락을 거부하고 생성 직후 운영 화면에서 조회된다."],
            ["FR-003\nMust", "판매 기록", "운영자가 +1·+5·+10을 누르면 판매수량·매출과 현재재고를 같은 처리 단위로 갱신한다.", "수량과 매출이 일치하고 재고가 0 미만이 되지 않는다."],
            ["FR-004\nMust", "재고 조정", "운영자가 재고 -10 또는 추가 준비 +10을 선택하면 새 Inventory Snapshot을 남긴다.", "변경량·시각·부스가 기록되고 이후 예측이 최신 재고를 읽는다."],
            ["FR-005\nMust", "통합 관제", "학생회가 전체 분석을 실행하면 부스별 판매·재고·30분 예측·예상 잔여·위험도를 비교한다.", "모든 활성 부스가 한 표에 표시되고 위험도가 텍스트와 색상으로 병기된다."],
            ["FR-006\nMust", "수요 예측", "10개 Feature가 준비되면 모델이 30분·60분·종료까지 판매량과 예상 잔여를 생성한다.", "Prediction에 timestamp, model_source, 근거와 Feature Snapshot이 연결된다."],
            ["FR-007\nMust", "폐기위험 판정", "예상 잔여를 현재 재고로 나눈 비율에 따라 LOW·MEDIUM·HIGH를 결정한다.", "9% LOW, 10% MEDIUM, 30% HIGH 경계 테스트가 통과한다."],
            ["FR-008\nMust", "예측 근거 표시", "최근 판매 추세, 강수확률, 공연 종료, 적용 할인 정보를 근거 목록으로 제공한다.", "운영 화면에서 최소 3개 근거와 모델 출처를 확인할 수 있다."],
        ],
        widths=[0.85, 1.25, 2.4, 1.3],
        font_size=7.35,
    )
    add_heading(doc, "공통 입력 검증", 2)
    make_key_value_table(
        doc,
        [
            ("식별자", "booth_id·menu_id·action_id는 존재 여부와 소유 관계를 확인하고 타 부스 자원 참조를 거부한다."),
            ("수량·가격", "판매·재고·가격은 정수 범위와 0 이상 조건을 확인하며, 화면 버튼과 서버 검증을 함께 적용한다."),
            ("시간", "Prediction과 Action은 ISO timestamp를 저장하고, 최신 Snapshot 기준으로 처리한다."),
        ],
        font_size=8.0,
    )

    page_break(doc)
    add_heading(doc, "기능 요구사항 상세 2", 2)
    make_table(
        doc,
        ["ID·우선순위", "요구사항", "선행 조건과 처리 결과", "수용 기준"],
        [
            ["FR-009\nMust", "Action 후보", "위험도와 운영조건으로 조리중단·할인·재고이동·홍보 후보와 사유를 생성한다.", "LOW는 유지, MEDIUM은 모니터링, HIGH는 정책 범위 후보만 생성한다."],
            ["FR-010\nMust", "사람 승인", "운영자가 PENDING 후보를 승인하면 동일 부스·유효 상태를 확인한 뒤 실행한다.", "승인 ID가 없거나 무효·타 부스·중복이면 실행 0건이다."],
            ["FR-011\nMust", "할인 Promotion", "20% 할인 Action이 실행되면 1시간 Promotion과 판매가를 생성한다.", "6,000원 상품이 4,800원으로 계산되고 승인 전에는 노출되지 않는다."],
            ["FR-012\nMust", "학생 혜택 노출", "학생 화면은 ACTIVE Promotion만 읽어 원가·할인가·부스·구역 정보를 보여준다.", "미승인 후보와 만료 혜택은 학생 화면에서 보이지 않는다."],
            ["FR-013\nMust", "반응 후 재예측", "판매·재고·할인 또는 학생 반응 상태가 변하면 Workflow를 다시 실행한다.", "새 Prediction이 저장되고 Before/After 잔여·위험도 차이를 확인한다."],
            ["FR-014\nShould", "Grounded AI 설명", "질문과 최신 DB Snapshot을 결합해 위험 원인·추천·승인 필요성을 설명한다.", "숫자는 저장된 값만 인용하고 데이터가 없으면 추측하지 않는다."],
            ["FR-015\nMust", "데이터 품질 검사", "필수 열·Null·중복·음수 Target·행 수·연도·메뉴 분포를 검사한다.", "오류 수와 quality score가 AI Ops에 표시되고 오류 데이터는 학습하지 않는다."],
            ["FR-016\nMust", "모델 학습·평가", "2023–2024 학습, 2025 검증으로 GBR을 평가하고 MAE·R²를 기록한다.", "고정 seed로 재현 가능하고 검증 연도와 행 수가 Run에 저장된다."],
            ["FR-017\nMust", "Registry·Audit", "Data Source, Model Run, Prediction, Agent Action 이력을 별도 화면에서 조회한다.", "ACTIVE 모델 1개와 입력→예측→승인→실행 추적이 가능하다."],
            ["FR-018\nMust", "Reset·Fallback", "Demo Reset으로 초기 상태를 복원하고 외부 의존성 장애 시 대체 경로를 사용한다.", "API 키와 네트워크 없이 핵심 폐루프 및 자동 테스트가 완료된다."],
        ],
        widths=[0.85, 1.25, 2.4, 1.3],
        font_size=7.1,
    )

    page_break(doc)
    add_heading(doc, "비기능 요구사항", 2)
    make_table(
        doc,
        ["ID", "품질 속성", "정량·정성 기준", "검증 방법·현재 범위"],
        [
            ["NFR-001", "사용성", "현장 핵심 입력 3회 이내, 주요 버튼 최소 56px, 실행 결과 즉시 피드백", "운영자 태스크 테스트·UI 점검 / MVP 반영"],
            ["NFR-002", "접근성", "색상+텍스트 이중부호, 키보드 포커스, Reduced Motion, 이미지 대체텍스트", "접근성 체크리스트·문서 Audit / 일부 반영"],
            ["NFR-003", "응답성", "단일 부스 추론 3초 이내, 전체 4부스 분석 5초 이내 목표", "파일럿 p95 측정 / 목표 기준"],
            ["NFR-004", "가용성", "Weather·LLM·LangGraph 장애에도 판매·예측·승인 핵심 흐름 유지", "의존성 제거 테스트 / MVP 반영"],
            ["NFR-005", "무결성", "판매와 재고 갱신 원자성, Foreign Key, 재고 음수·중복 실행 방지", "DB 제약·트랜잭션 테스트 / 보강 필요"],
            ["NFR-006", "보안", "API 키는 환경변수, 불필요한 개인정보 0개, 역할별 최소 권한", "Secret Scan·스키마 검토 / 인증은 파일럿 추가"],
            ["NFR-007", "설명가능성", "모델 출처, 입력 근거 3개 이상, 위험 임계값, Action 사유 표시", "화면·Prediction Audit / MVP 반영"],
            ["NFR-008", "감사가능성", "입력·예측·승인·실행·재예측에 timestamp와 ID 저장", "Audit 조회·추적성 표 / MVP 반영"],
            ["NFR-009", "재현성", "seed 42, 모델 파라미터·학습/검증 기간·지표·Run ID 보존", "동일 데이터 재학습 비교 / MVP 반영"],
            ["NFR-010", "확장성", "캠퍼스·부스·메뉴 키로 분리하고 Adapter로 CSV·API·POS 교체", "3개 대학 이력 계약·Adapter 검토 / 구조 반영"],
            ["NFR-011", "유지보수성", "UI·서비스·모델·Agent·DB 계층 분리와 자동 테스트", "모듈 의존성·pytest / 7건 통과"],
            ["NFR-012", "운영 안전", "무승인 실행 0건, 이전 모델 Rollback, 정책 밖 할인 차단", "Agent·Model Gate / 일부 파일럿 추가"],
        ],
        widths=[0.75, 1.0, 2.65, 1.4],
        font_size=7.2,
    )
    add_body(
        doc,
        "응답시간·동시사용자·가용률은 실제 운영 환경에서 측정할 목표값이다. MVP 시연 결과를 운영 SLA로 표현하지 않으며, "
        "파일럿에서 p95 지연, 오류율, 입력 누락률을 수집해 기준을 확정한다.",
        size=8.5,
    )

    page_break(doc)
    add_heading(doc, "역할 권한과 상태 모델", 2)
    add_heading(doc, "역할 기반 권한 매트릭스", 2)
    make_table(
        doc,
        ["기능·데이터", "학생회 관리자", "부스 운영자", "학생", "AI 운영자"],
        [
            ["전체 부스 현황·예측", "조회·전체 분석", "자기 부스 조회", "공개 혜택만", "Audit 조회"],
            ["판매·재고 변경", "조회", "자기 부스 입력", "불가", "원본 변경 불가"],
            ["Action 승인·실행", "정책 확인", "자기 부스 승인", "불가", "이력 조회"],
            ["Promotion", "전체 조회", "자기 부스 생성 승인", "ACTIVE 조회", "Audit 조회"],
            ["데이터·모델 Pipeline", "상태 조회", "불가", "불가", "검증·학습·등록"],
            ["Demo Reset", "허용", "불가", "불가", "허용"],
        ],
        widths=[1.45, 1.2, 1.2, 0.95, 1.0],
        font_size=7.65,
    )
    add_body(
        doc,
        "MVP는 역할별 Route 분리로 권한 경계를 시연한다. 실서비스에서는 대학 계정 SSO, 역할 Claim, 부스 소유권 검증, "
        "관리자 승인 정책을 서버 계층에 추가해야 한다.",
        size=8.55,
    )
    add_heading(doc, "핵심 상태 전이", 2)
    make_table(
        doc,
        ["객체", "상태 흐름", "전이 조건", "금지 조건"],
        [
            ["Agent Action", "PENDING → APPROVED → EXECUTED", "동일 부스 운영자가 유효 Action ID 승인", "타 부스·이미 처리·없는 ID 전이 금지"],
            ["Promotion", "없음 → ACTIVE → EXPIRED", "DISCOUNT 실행과 유효 시간", "승인 전·만료 후 학생 노출 금지"],
            ["Model Run", "CANDIDATE → REVIEWED → ACTIVE → ARCHIVED", "품질·성능 Gate와 사람 검토", "검증 실패 모델 활성화 금지"],
            ["Data Source", "INGESTED → VALIDATED → READY 또는 QUARANTINED", "필수 열·값·출처 검증", "오류 데이터 학습 금지"],
            ["Prediction", "Snapshot 생성 → 비교 → 보존", "입력 또는 상태 변경", "과거 Snapshot 덮어쓰기 금지"],
        ],
        widths=[1.1, 1.8, 1.8, 1.1],
        font_size=7.65,
    )

    page_break(doc)
    add_heading(doc, "서비스 인터페이스와 예외 흐름", 2)
    add_heading(doc, "MVP 서비스 계약과 운영 API 확장", 2)
    make_table(
        doc,
        ["업무", "현재 MVP 함수", "입력 → 출력", "실서비스 API 예시"],
        [
            ["판매 기록", "record_sale", "booth_id, quantity → Sales·Inventory Snapshot", "POST /booths/{id}/sales"],
            ["재고 조정", "adjust_stock", "booth_id, delta → Inventory Snapshot", "PATCH /booths/{id}/inventory"],
            ["분석 실행", "run_workflow", "booth_id → Prediction·Candidates", "POST /booths/{id}/analyses"],
            ["Action 승인", "run_workflow + action_id", "approved_action_ids → Executed Action", "POST /actions/{id}/approve"],
            ["학생 혜택", "get_active_promotions", "현재 시각 → ACTIVE Promotion", "GET /promotions?status=active"],
            ["AI 설명", "ask_admin", "question + DB Snapshot → answer, engine", "POST /assistant/explanations"],
            ["모델 학습", "run_training_pipeline", "Historical Data → Run·Metric·Registry", "POST /ml/runs"],
        ],
        widths=[1.0, 1.45, 2.15, 1.2],
        font_size=7.35,
    )
    add_body(
        doc,
        "현재 Streamlit MVP는 Python 서비스 함수를 직접 호출한다. API 경로는 파일럿 확장 계약이며, 인증·멱등키·요청 버전·오류 코드를 "
        "추가한 뒤 프런트엔드와 서비스 계층을 분리한다.",
        size=8.4,
    )
    add_heading(doc, "주요 예외와 복구", 2)
    make_table(
        doc,
        ["ID", "예외 상황", "사용자 피드백·시스템 처리", "상태"],
        [
            ["EX-01", "부스·메뉴 없음", "처리를 중단하고 입력 대상을 다시 선택하게 한다.", "보강 필요"],
            ["EX-02", "판매량 오류·재고 부족", "유효 범위를 안내하고 DB 쓰기를 수행하지 않는다.", "보강 필요"],
            ["EX-03", "Weather API 실패", "Mock Weather 출처를 표시하고 예측을 계속한다.", "구현"],
            ["EX-04", "ML 추론 실패", "Transparent heuristic fallback과 모델 출처를 표시한다.", "구현"],
            ["EX-05", "LLM 키·응답 없음", "Local Grounded Chat으로 같은 DB 근거를 설명한다.", "구현"],
            ["EX-06", "LangGraph 의존성 실패", "동일 Node 순서의 Sequential fallback을 실행한다.", "구현"],
            ["EX-07", "Action ID 무효·중복", "실행하지 않고 기존 상태와 Audit을 유지한다.", "구현"],
            ["EX-08", "동시 재고 갱신 충돌", "버전 비교 후 재조회·재시도를 요구한다.", "파일럿 추가"],
        ],
        widths=[0.65, 1.55, 2.85, 0.75],
        font_size=7.25,
    )

    page_break(doc)
    add_heading(doc, "요구사항 추적성과 시험 계획", 2)
    make_table(
        doc,
        ["요구사항", "화면·모듈", "저장 데이터", "시험·증거"],
        [
            ["FR-002", "Operator·database", "booths, menus, inventory", "pytest: 부스 등록 최소 필드"],
            ["FR-003·004", "Operator·database", "sales_snapshots, inventory_snapshots", "수량·재고 UI 통합 시험"],
            ["FR-006·007", "demand_model·waste_risk", "predictions", "pytest: 위험 임계값 경계"],
            ["FR-009·010", "rules·LangGraph", "agent_actions", "pytest: 할인 승인 폐루프"],
            ["FR-011·012·013", "Operator·Student·Simulation", "promotions, predictions", "할인가 4,800원·70→15 재예측 시험"],
            ["FR-014", "Admin Chat·chat", "prediction·action snapshot", "pytest: 현재 예측 기반 Chat"],
            ["FR-015·016·017", "AI Ops·mlops", "data_sources, ml_runs", "pytest: ACTIVE 모델 등록"],
            ["FR-018", "Gateway·Graph·Adapters", "settings·Audit", "키·네트워크 제거 복구 시험"],
        ],
        widths=[0.85, 1.55, 1.45, 1.95],
        font_size=7.45,
    )
    add_heading(doc, "인수 시험 시나리오", 2)
    make_table(
        doc,
        ["TC", "Given", "When", "Then"],
        [
            ["TC-01", "닭꼬치 재고 180·할인 없음", "전체 분석", "예상 잔여 70·HIGH와 근거 표시"],
            ["TC-02", "할인 Action PENDING", "승인 전 학생 화면", "할인가 미노출"],
            ["TC-03", "유효 DISCOUNT Action", "운영자 20% 승인", "EXECUTED와 4,800원 Promotion 생성"],
            ["TC-04", "Promotion ACTIVE", "학생 화면 조회", "부스·구역·원가·할인가 일치"],
            ["TC-05", "학생 반응 시뮬레이션", "재예측", "예상 잔여 15·LOW와 새 Snapshot"],
            ["TC-06", "OPENAI_API_KEY 없음", "할인 이유 질문", "70개·승인 필요성을 Local Chat이 설명"],
            ["TC-07", "2023–2025 합성 이력", "Pipeline 실행", "품질 100%, 2025 Hold-out, ACTIVE Run"],
            ["TC-08", "타 부스·중복 Action ID", "실행 요청", "실행 0건·기존 상태 유지"],
        ],
        widths=[0.65, 1.7, 1.5, 1.95],
        font_size=7.25,
    )
    add_body(
        doc,
        "완료 정의: Must 요구사항과 자동 테스트가 통과하고, 합성 데이터·모델 출처·승인 경계가 화면에 표시되며, "
        "Gateway→Admin→Operator→Student→재예측→AI Ops 시연이 Reset 후 반복 가능해야 한다.",
        size=8.5,
    )

    # Requirements 8-2
    page_break(doc)
    add_heading(doc, "(8-2) 주요 화면 및 사용자 흐름", 2)
    make_table(
        doc,
        ["순서", "화면", "사용자 행동", "시스템 상태 변화"],
        [
            ["1", "학생회 관제", "닭꼬치 재고 180, 예상 잔여 70, HIGH 확인", "예측 스냅샷과 근거 표시"],
            ["2", "학생회 관제", "왜 할인해야 하는지 AI Chat 질문", "DB 스냅샷에 근거한 설명"],
            ["3", "부스 운영자", "20% 마감 할인 승인", "Action EXECUTED, Promotion 활성화"],
            ["4", "학생", "6,000원 → 4,800원 혜택 확인", "같은 SQLite 상태를 즉시 반영"],
            ["5", "Before/After", "학생 반응 발생 후 재예측", "예상 잔여 15, LOW로 갱신"],
            ["6", "AI Ops", "Pipeline·Model·Agent 이력 확인", "품질·성능·승인 감사 증거 제시"],
        ],
        widths=[0.55, 1.25, 2.35, 1.65],
        font_size=8.2,
    )
    add_heading(doc, "휴먼 인터페이스 요구사항", 2)
    make_key_value_table(
        doc,
        [
            ("역할 격리", "학생회·운영자·학생·AI 운영자의 과업과 정보 밀도를 분리해 선택 부담을 줄인다."),
            ("조작 접근성", "현장 버튼은 최소 56px, 주요 Action은 상→하 흐름, 키보드 Focus Ring과 Reduced Motion을 지원한다."),
            ("상태 인지", "위험은 색상만 사용하지 않고 점·HIGH/MEDIUM/LOW 텍스트·설명을 함께 표시한다."),
            ("오류 예방", "가격·조리·이동은 AI가 자동 실행하지 않으며 승인 버튼 가까이에 근거와 영향 범위를 표시한다."),
            ("피드백", "승인 직후 학생 가격, 학생 반응 직후 Before/After와 Dashboard가 같은 상태를 읽어 즉시 변화한다."),
            ("신뢰", "합성 데이터 경고, 예측 근거, 모델 출처, LLM 역할 제한을 행동 지점 가까이에 배치한다."),
        ],
        font_size=8.65,
    )

    # Requirements 8-3
    page_break(doc)
    add_heading(doc, "(8-3) AI 및 데이터 구현 계획", 2)
    make_key_value_table(
        doc,
        [
            ("AI 모델", "GradientBoostingRegressor로 30분 판매량을 예측하고 종료까지 확장한다. LLM은 예측 숫자를 만들지 않고 DB 스냅샷을 설명한다."),
            ("데이터", "University, Festival, Booth, Menu, InventorySnapshot, SalesSnapshot, Weather, FestivalEvent, Prediction, AgentAction, Promotion 스키마. 기본 이력은 합성 CSV 216행."),
            ("외부 서비스", "Open-Meteo Adapter는 선택 사항이며 실패 시 Mock Weather로 복귀한다. OpenAI API는 선택형 설명 계층이고 키가 없으면 Local Grounded Chat을 사용한다."),
            ("구현 도구", "Python, Pandas, scikit-learn, Streamlit, Plotly, LangGraph, SQLite, pytest, python-dotenv."),
            ("운영 관리", "AI Ops에서 Ingest·Validate·Feature·Train·Evaluate·Register, Data Registry, Model Registry, Prediction Audit, Action Queue를 분리 관리한다."),
        ],
        font_size=8.9,
    )
    add_heading(doc, "파이프라인 입출력과 통제", 2)
    make_table(
        doc,
        ["단계", "입력", "출력", "품질·안전 통제"],
        [
            ["Ingest", "과거 CSV·당일 SQLite", "표준 스키마", "버전·출처·data_class 기록"],
            ["Validate", "필수 열·Target", "품질 점수", "Null·중복·음수 Target 차단"],
            ["Feature", "판매·시간·환경", "10개 feature", "학습·추론 동일 계약"],
            ["Train", "2023–2024 합성 이력", "GBR 모델", "고정 seed·파라미터 기록"],
            ["Evaluate", "2025 hold-out", "MAE·R²", "시간 순서 분리·합성 경고"],
            ["Register", "모델·지표·검토", "ACTIVE run", "사람 검토 후 활성화"],
            ["Infer", "현재 부스 상태", "예측·위험·근거", "모델 출처와 fallback 표시"],
            ["Act", "후보·승인 ID", "Action·Promotion", "운영자 승인·감사 로그"],
        ],
        widths=[0.8, 1.45, 1.3, 2.25],
        font_size=7.75,
    )

    # Requirements 8-4: detailed data design
    page_break(doc)
    add_heading(doc, "(8-4) 데이터 상세 설계", 2)
    add_body(
        doc,
        "ZeroFest AI는 과거 축제 데이터만으로 결과를 고정하지 않는다. 과거 이력은 시간대·메뉴·캠퍼스 규모에 따른 "
        "기본 수요 패턴을 학습하는 데 사용하고, 행사 당일에는 부스별 최근 판매, 직전 판매, 현재 재고, 날씨, 공연 종료, "
        "승인된 할인율을 결합해 판매 입력 또는 상태 변화 때마다 다시 추론한다. MVP의 2023–2025 데이터 216행은 "
        "모두 SAMPLE / SYNTHETIC이며, 실제 운영 전 대학별 실데이터로 동일 계약을 검증해야 한다.",
        size=9.1,
    )
    add_heading(doc, "데이터 원천과 갱신 주기", 2)
    make_table(
        doc,
        ["데이터 원천", "핵심 필드", "갱신·단위", "사용 목적과 통제"],
        [
            ["과거 축제 이력", "연도, 대학, 학생수, 메뉴, 가격, 판매, 날씨, 공연, 할인", "행사 종료 후 적재\n30분 단위", "모델 학습·시간 Hold-out. data_class와 출처를 기록한다."],
            ["당일 판매", "timestamp, booth_id, menu_id, sales_quantity, revenue", "입력 즉시\n30분 집계", "최근·직전 판매량과 추세를 생성한다. 음수·중복 입력을 차단한다."],
            ["당일 재고", "current_stock, additional_stock", "판매·보충 시 즉시", "종료 예상 잔여의 기준값. 판매와 재고를 한 트랜잭션으로 갱신한다."],
            ["날씨", "기온, 강수확률, 강우량, 습도, provider", "최대 30분", "1시간 예보를 특징값으로 사용한다. API 실패 시 버전 관리 Mock을 사용한다."],
            ["공연 일정", "구역, 시작·종료, 예상 관객", "일정 변경 시", "1시간 이내 공연 종료 여부와 구역 수요 변화를 계산한다."],
            ["운영 결과", "예측, 위험도, 후보, 승인·실행, Promotion", "워크플로 단계별", "입력→예측→승인→실행→재예측을 timestamp와 ID로 추적한다."],
        ],
        widths=[1.05, 2.0, 1.0, 1.75],
        font_size=7.65,
    )
    add_heading(doc, "학습 데이터 계약", 2)
    make_key_value_table(
        doc,
        [
            ("예측 단위", "대학축제 × 부스 × 메뉴 × 30분 시점. Target은 다음 30분 판매량 future_sales_30m이다."),
            ("시간 분리", "2023–2024년 144행으로 학습하고 2025년 72행을 시간 Hold-out으로 검증한다. 검증 통과 후 전체 2023–2025 이력으로 운영 모델을 재학습한다."),
            ("누수 방지", "현재 시점 이후의 판매·재고와 미래 행사 결과는 특징값에 포함하지 않는다. 할인율은 해당 시점에 이미 승인된 값만 사용한다."),
            ("실데이터 전환", "HISTORICAL_DATA_PATH Adapter에 동일 필수 열을 가진 CSV를 연결하고, 대학·메뉴별 오차 검토 후 활성화한다."),
        ],
        font_size=8.15,
    )

    page_break(doc)
    add_heading(doc, "특징값과 예측 산출 기준", 2)
    make_table(
        doc,
        ["Feature", "정의·형식", "생성 시점", "예측에서의 의미"],
        [
            ["recent_sales_30m", "최근 30분 판매 수량 / 정수", "당일 판매 집계", "가장 가까운 현재 수요 수준"],
            ["previous_sales_30m", "직전 30분 판매 수량 / 정수", "당일 판매 집계", "증가·감소 추세 비교 기준"],
            ["hour", "현재 시각 16–22 / 정수", "추론 시각", "시간대별 유동 인구 패턴"],
            ["minutes_to_close", "축제 종료까지 분 / 정수", "현재 시각과 종료 시각", "남은 판매 기회"],
            ["precipitation_probability", "1시간 뒤 강수확률 0–100 / 실수", "Weather Adapter", "야외 체류와 구매 감소 가능성"],
            ["event_ending_soon", "1시간 이내 공연 종료 0·1", "공연 일정", "관객 이동에 따른 단기 수요 변화"],
            ["discount_rate", "승인된 할인율 0–100 / 정수", "Promotion 상태", "가격 개입 후 수요 반응"],
            ["category_code", "food 0, drink 1, dessert 2", "메뉴 Master", "메뉴군별 판매 패턴"],
            ["price_k", "판매가 ÷ 1,000 / 실수", "메뉴 Master", "가격 수준에 따른 구매 차이"],
            ["campus_scale", "재학생 수 ÷ 10,000 / 실수", "대학 Master", "대학 규모 정규화"],
        ],
        widths=[1.25, 1.8, 1.35, 1.4],
        font_size=7.4,
    )
    add_heading(doc, "예측값과 위험도 계산", 2)
    make_table(
        doc,
        ["산출값", "계산 및 의미", "운영 사용"],
        [
            ["30분 판매 예측", "GradientBoostingRegressor의 직접 출력. 최소 1개로 제한한다.", "가장 가까운 수요와 근거 표시"],
            ["60분 판매 예측", "MVP에서 30분 예측 × 1.86으로 확장한다.", "재고·조리 준비의 단기 기준"],
            ["종료까지 판매 예측", "30분 예측 × 남은 30분 구간 × 0.82, 현재 재고 이내로 제한한다.", "종료 예상 잔여 계산"],
            ["예상 잔여", "현재 재고 − 종료까지 판매 예측, 최소 0개", "폐기 가능 수량의 대리 지표"],
            ["폐기위험", "예상 잔여 ÷ 현재 재고. LOW <10%, MEDIUM 10–<30%, HIGH ≥30%", "결정론적 규칙으로 Action 후보 생성"],
        ],
        widths=[1.15, 3.15, 1.5],
        font_size=7.85,
    )
    add_body(
        doc,
        "고정 시연의 70개 HIGH → 15개 LOW 결과는 발표 재현성을 위한 Demo Simulation 보정이며 실제 효과 검증값이 아니다. "
        "현장 배포에서는 60분·종료 예측도 별도 Target으로 학습하거나 시계열 모델과 비교 검증한다.",
        size=8.5,
    )

    # Requirements 8-5: detailed ML lifecycle
    page_break(doc)
    add_heading(doc, "(8-5) 머신러닝 학습 및 실시간 추론", 2)
    add_figure(
        doc,
        assets["offline_online_workflow"],
        "그림 8. 과거 축제 학습과 당일 운영 신호를 결합하는 오프라인·온라인 AI 워크플로",
        width=5.78,
    )
    add_heading(doc, "모델 선택과 학습 설정", 2)
    make_table(
        doc,
        ["항목", "MVP 설정", "선택 이유와 확장 기준"],
        [
            ["알고리즘", "GradientBoostingRegressor", "소규모 표형 데이터와 비선형 상호작용에 적합하고 Feature Importance를 제공한다."],
            ["하이퍼파라미터", "120 trees, depth 3, learning rate 0.045, Huber loss, seed 42", "이상치 영향을 줄이고 시연 재현성을 확보한다. 튜닝 결과로 오인하지 않는다."],
            ["검증 방식", "2025 시간 Hold-out", "무작위 분할로 미래 정보가 학습에 섞이는 문제를 줄인다."],
            ["현재 합성 지표", "MAE 1.94개/30분, R² 0.877", "합성 데이터 내부 결과이며 실제 대학 성능을 의미하지 않는다."],
            ["운영 추론", "판매·재고 변화 즉시 또는 30분 주기", "동일 Feature 계약과 모델 출처를 Prediction Snapshot에 저장한다."],
        ],
        widths=[1.15, 2.05, 2.6],
        font_size=7.85,
    )

    # Requirements 8-6: agent and MLOps controls
    page_break(doc)
    add_heading(doc, "(8-6) AI Agent 상세 워크플로와 통제", 2)
    make_table(
        doc,
        ["순서·Node", "읽는 상태", "생성·변경 상태", "오류·안전 처리"],
        [
            ["1 상태 로드", "booth_id, SQLite", "부스·재고·최근 판매, 현재시각", "존재하지 않는 부스는 실행 중단"],
            ["2 맥락 결합", "Weather, Event", "1시간 날씨·공연 종료 맥락", "API 실패 시 Mock Weather"],
            ["3 수요 예측", "10개 Feature", "30분·60분·종료 판매, 잔여, 근거", "모델 실패 시 투명한 가중 휴리스틱"],
            ["4 위험 계산", "예상 잔여·현재 재고", "LOW·MEDIUM·HIGH, 잔여비율", "고정 임계값으로 감사 가능성 유지"],
            ["5 운영조건", "판매 추세·비·공연", "감소, 강우, 공연 종료 여부", "조건과 예측을 분리 저장"],
            ["6 후보 생성", "위험도·운영조건", "조리중단·할인·이동·홍보 후보", "LOW는 유지, 정책 밖 할인 생성 금지"],
            ["7 사람 승인", "후보·approved_action_ids", "PENDING 또는 승인 분기", "승인 ID가 없으면 실행 없이 종료"],
            ["8 Action 실행", "유효한 PENDING Action", "APPROVED→EXECUTED, Promotion", "타 부스·무효·중복 ID는 건너뜀"],
            ["9 모니터·재예측", "변경된 판매·재고·할인", "새 Prediction과 위험도", "동일 입력·모델 출처·결과를 Audit"],
        ],
        widths=[1.05, 1.45, 1.75, 1.55],
        font_size=7.25,
    )
    add_heading(doc, "AI 구성요소별 책임 경계", 2)
    make_key_value_table(
        doc,
        [
            ("Machine Learning", "판매 수량만 예측한다. 위험 판정, 할인 결정, 실행 권한을 갖지 않는다."),
            ("결정 규칙", "잔여비율 임계값과 Action 정책을 코드로 고정해 모델이 바뀌어도 운영 기준을 추적할 수 있게 한다."),
            ("LangGraph", "상태 전달과 승인 분기를 관리한다. 의존성 장애 시 같은 Node 순서의 Sequential fallback을 사용한다."),
            ("LLM 설명", "DB의 최신 예측·근거·Action 상태를 자연어로 설명한다. 데이터를 찾지 못하면 추측하지 않고 제한을 알린다."),
            ("운영자", "할인·조리중단·재고 이동의 최종 결정자다. 승인·실행 시각과 Action ID를 남긴다."),
        ],
        font_size=8.0,
    )

    page_break(doc)
    add_heading(doc, "AI Ops 운영과 모델 수명주기", 2)
    make_table(
        doc,
        ["대시보드 영역", "관리 정보", "판단 기준", "조치"],
        [
            ["Data Registry", "출처, 기간, 행 수, data_class, 검증 시각", "필수 열 누락 0, Null·중복·음수 Target 0", "오류 데이터 격리 후 재적재"],
            ["Pipeline Control", "Ingest→Validate→Feature→Train→Evaluate→Register", "모든 단계 PASS, Register는 사람 검토", "실패 단계부터 재실행"],
            ["Model Registry", "run_id, 학습 연도, 검증 연도, MAE, R², ACTIVE", "기준 모델 대비 오차·캠퍼스별 편차 확인", "승인·활성화 또는 이전 모델 유지"],
            ["Prediction Audit", "시각, 부스, 예측량, 잔여, 위험도, model_source", "예측과 실제 30분 판매를 연결 가능", "오차가 큰 시간대·메뉴 분석"],
            ["Agent Audit", "후보, 사유, 상태, 승인·실행 시각", "무승인 실행 0건, 중복 실행 0건", "정책 위반 Action 차단·조사"],
            ["Fallback 상태", "Weather·LLM·LangGraph 사용 경로", "대체 경로에서도 핵심 Demo 완료", "장애 원인 기록 후 정상 경로 복구"],
        ],
        widths=[1.15, 1.75, 1.65, 1.25],
        font_size=7.5,
    )
    add_heading(doc, "재학습과 배포 기준", 2)
    make_key_value_table(
        doc,
        [
            ("재학습 시점", "축제 종료 후 실제 30분 판매 Target이 확정되거나 충분한 신규 라벨 구간이 쌓일 때 Batch 재학습한다. 행사 중에는 기본적으로 추론만 수행한다."),
            ("승격 Gate", "기존 ACTIVE 모델보다 전체 MAE가 개선되고 대학·메뉴별 성능 저하가 허용 범위 이내이며 Data·Model·Human Gate를 모두 통과해야 한다."),
            ("Drift 대응", "실서비스에서는 최근 실제값 대비 Rolling MAE와 특징값 분포를 감시한다. 기준 초과 시 경고, 이전 모델 유지, 규칙 기반 운영으로 단계적으로 축소한다."),
            ("Rollback", "모델 run_id와 Feature 계약을 보존하고, 승인된 이전 ACTIVE 모델로 되돌린다. 이미 실행된 Action은 별도 감사 기록으로 유지한다."),
            ("현재 구현 범위", "MVP는 품질 점수, 시간 Hold-out, Model Registry, Prediction·Agent Audit을 구현했다. 자동 Drift 판정과 무중단 배포는 파일럿 단계에서 추가한다."),
        ],
        font_size=8.0,
    )

    # Safety
    page_break(doc)
    add_heading(doc, "9. 안전·윤리·사용성 점검", 1)
    make_key_value_table(
        doc,
        [
            ("데이터 및 개인정보", "■ 주민등록번호·전화번호·주소를 수집하지 않는다.\n■ 기본 데이터는 가상 대학 합성 이력이며 화면과 문서에 표시한다.\n■ 정밀 GPS 대신 익명 Zone 수준 정보만 확장 후보로 둔다.\n■ API 키는 .env에만 저장하고 코드·화면·로그에 노출하지 않는다.\n■ 실제 데이터 도입 시 목적·보유기간·삭제 절차와 접근권한을 정의한다."),
            ("AI 안전·윤리", "■ 예측을 사실이나 확정값으로 단정하지 않고 근거·모델 출처·한계를 표시한다.\n■ 할인·조리중단·이동은 운영자 승인 전 실행하지 않는다.\n■ 메뉴·부스·캠퍼스별 오차를 분리 점검해 특정 집단에 지속적으로 불리한 추천을 찾는다.\n■ LLM은 설명 전용이며 수치 예측과 실행 권한을 갖지 않는다.\n■ 승인·거절·실행·재예측을 감사 가능하게 기록한다."),
            ("사용성 및 접근성", "■ 역할별 화면으로 정보 과부하를 줄인다.\n■ 버튼은 큰 터치 영역과 명확한 동사를 사용한다.\n■ 색상과 함께 텍스트·아이콘·수치를 병기한다.\n■ 오류·fallback·데이터 범위를 화면에 알린다.\n■ 한 번 클릭 Reset과 샘플 시나리오로 처음부터 끝까지 재현한다.\n■ 키보드 포커스와 Reduced Motion을 고려한다."),
        ],
        label_width=1.35,
        value_width=4.45,
        font_size=8.55,
    )
    add_heading(doc, "운영 전 필수 게이트", 2)
    make_table(
        doc,
        ["Gate", "확인 항목", "통과 기준"],
        [
            ["Data", "출처·동의·품질·보유기간", "미확인 데이터 학습 금지"],
            ["Model", "캠퍼스·메뉴별 오차와 drift", "기준 초과 시 이전 모델 또는 규칙 fallback"],
            ["Action", "가격 하한·시간창·재고 임계값", "정책 밖 후보 생성·실행 금지"],
            ["Human", "운영자 승인과 취소 경로", "무승인 실행 0건"],
            ["Audit", "입력·예측·승인·실행·재예측", "타임스탬프와 출처 추적 가능"],
        ],
        widths=[0.8, 2.7, 2.3],
        font_size=8.25,
    )

    # Judging alignment
    page_break(doc)
    add_heading(doc, "10. 심사 기준 대응 및 발표 전략", 1)
    make_table(
        doc,
        ["심사 기준", "제출물의 핵심 근거", "발표에서 보여줄 증거"],
        [
            ["기획성 및 주제 적합성 25", "폐기 전 예방이라는 명확한 문제, 4역할 가치, KPI와 구매·사용·가격·도입 시장검증 계획", "문제 연쇄·시장 Gate·폐루프 차별성 40초"],
            ["AI 및 데이터 기술 활용도 25", "시간 Hold-out ML, 위험 규칙, LangGraph, Grounded LLM, AI Ops", "모델 출처·근거·승인 분리 50초"],
            ["기술적 구체성 및 구현 가능성 25", "SQLite 단일 상태, 6단계 파이프라인, fallback, 7개 자동 테스트", "Admin→Operator→Student→재예측 2분"],
            ["발표 전달력 및 질의응답 25", "고정 3–5분 시나리오, Before/After, 한계의 선제적 공개", "AI Ops와 예상 Q&A 40초"],
        ],
        widths=[1.6, 2.55, 1.65],
        font_size=8.15,
    )
    add_heading(doc, "3분 30초 시연 흐름", 2)
    make_table(
        doc,
        ["시간", "화면", "메시지"],
        [
            ["0:00–0:40", "게이트웨이·Admin", "남은 음식을 처리하는 것이 아니라 남지 않도록 운영을 바꾼다."],
            ["0:40–1:30", "학생회 관제", "재고 180, 예상 잔여 70 HIGH와 비·공연 종료 근거를 확인한다."],
            ["1:30–2:15", "부스 운영자", "AI는 제안하고, 사람의 20% 할인 승인 뒤에만 실행된다."],
            ["2:15–2:50", "학생", "6,000원 → 4,800원이 같은 상태에서 즉시 노출된다."],
            ["2:50–3:30", "Before/After·AI Ops", "학생 반응 후 70 → 15 LOW 재예측과 Pipeline·Audit를 확인한다."],
        ],
        widths=[1.0, 1.45, 3.35],
        font_size=8.3,
    )
    add_heading(doc, "예상 질의응답", 2)
    make_table(
        doc,
        ["질문", "답변 요지"],
        [
            ["전년도 데이터로 예측하나요?", "과거 연도 이력으로 기본 패턴을 학습하되, 당일 최근 판매·재고·날씨·공연·할인 신호로 30분마다 보정합니다."],
            ["실제 데이터와 성과인가요?", "아닙니다. 216행과 70→15는 합성·시뮬레이션입니다. 실제 대학 파일럿에서 MAE와 폐기 kg을 별도 검증합니다."],
            ["왜 딥러닝이 아닌가요?", "현재 데이터 규모에서는 작은 표형 데이터에 적합하고 설명·재현이 쉬운 Gradient Boosting이 구현 가능성과 과적합 관리에 유리합니다."],
            ["예측이 틀리면 어떻게 하나요?", "근거와 오차를 표시하고, 어떤 Action도 운영자 승인 전 실행하지 않으며, 30분 rolling 재예측과 fallback을 사용합니다."],
            ["LLM은 무엇을 하나요?", "판매량을 만들지 않고 현재 DB 예측 스냅샷을 자연어로 설명합니다. API가 없어도 같은 근거의 로컬 설명이 작동합니다."],
            ["가장 큰 차별점은 무엇인가요?", "예측 대시보드가 아니라 운영 승인, 학생 혜택, 실제 반응, 재예측까지 닫힌 루프로 구현했다는 점입니다."],
            ["누가 비용을 지불하나요?", "학생회·대학 행사 담당 부서·축제 대행사를 구매자 가설로 두고 있습니다. 구매자 인터뷰, 행사 단위 가격 카드, 파일럿 제안과 도입의향서로 실제 지불 주체와 범위를 검증합니다."],
        ],
        widths=[1.65, 4.15],
        font_size=8.15,
    )
    add_body(doc, "최종 전달 원칙: 성능 숫자보다 의사결정 구조를 보여주고, 합성 데이터의 한계를 먼저 공개한 뒤 실제 파일럿에서 무엇을 측정할지 답한다.", bold_lead="최종 전달 원칙:", size=9.2)

    normalize_page_breaks(doc)
    configure_running_headers(doc)

    # Ensure section geometry remains the template geometry.
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.18)
    section.right_margin = Inches(1.18)
    section.top_margin = Inches(1.38)
    section.bottom_margin = Inches(0.59)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != source_sha:
        raise RuntimeError("Reference template was modified during generation.")
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
