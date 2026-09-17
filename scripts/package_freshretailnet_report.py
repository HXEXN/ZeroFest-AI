"""Package validated FreshRetailNet report pages into a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
OUTPUT = OUTPUT_DIR / "FreshRetailNet_상관분석_히트맵_보고서.docx"
PAGES = [OUTPUT_DIR / f"freshretail_report_page_{index}.png" for index in range(1, 4)]


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.12)
    section.bottom_margin = Inches(0.12)
    section.left_margin = Inches(0.15)
    section.right_margin = Inches(0.15)
    for index, path in enumerate(PAGES):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if index:
            paragraph.paragraph_format.page_break_before = True
        paragraph.add_run().add_picture(str(path), width=Inches(7.75))
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
