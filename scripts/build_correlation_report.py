"""Create a heatmap-based correlation analysis report for ZeroFest synthetic history."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
OUTPUT = OUTPUT_DIR / "ZeroFest_변수간_상관분석_보고서.docx"
HEATMAP = OUTPUT_DIR / "zerofest_correlation_heatmap.png"
HISTORY_PATH = ROOT / "data" / "sample_historical_sales.csv"
FEATURES = [
    "recent_sales_30m", "previous_sales_30m", "festival_day", "hour", "minute_of_day",
    "second_of_minute", "minutes_to_close", "seconds_to_close", "minutes_to_festival_end",
    "precipitation_probability", "temperature", "event_ending_soon", "discount_rate",
    "category_code", "price_k", "campus_scale",
]

LABELS = {
    "recent_sales_30m": "최근 30분 판매",
    "previous_sales_30m": "이전 30분 판매",
    "festival_day": "축제 일차",
    "hour": "시",
    "minute_of_day": "분 단위 시각",
    "second_of_minute": "초",
    "minutes_to_close": "당일 종료까지 분",
    "seconds_to_close": "당일 종료까지 초",
    "minutes_to_festival_end": "축제 전체 종료까지 분",
    "precipitation_probability": "강수확률",
    "temperature": "기온",
    "event_ending_soon": "공연 종료 임박",
    "discount_rate": "할인율",
    "category_code": "메뉴 유형",
    "price_k": "가격 천원",
    "campus_scale": "캠퍼스 규모",
    "future_sales_30m": "다음 30분 판매량",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "AppleGothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2:
                set_cell_shading(cells[index], "F3F7FA")
            set_cell_border(cells[index])
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_body(document: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.35
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.name = "AppleGothic"
        lead._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
    run = paragraph.add_run(text)
    run.font.name = "AppleGothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "AppleGothic"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")


def build_heatmap(correlation: pd.DataFrame) -> None:
    labels = list(correlation.columns)
    cell, left, top = 43, 270, 90
    width = left + cell * len(labels) + 70
    height = top + cell * len(labels) + 120
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 12)
    small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 9)
    title = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 22)
    draw.text((left, 25), "ZeroFest feature correlation heatmap", font=title, fill="black")

    def cell_color(value: float) -> tuple[int, int, int]:
        intensity = abs(value)
        base = (43, 98, 160) if value >= 0 else (190, 55, 65)
        return tuple(int(255 * (1 - intensity) + channel * intensity) for channel in base)

    for row, row_label in enumerate(labels):
        draw.text((8, top + row * cell + 14), row_label, font=regular, fill="black")
        for col, col_label in enumerate(labels):
            value = float(correlation.iat[row, col])
            x, y = left + col * cell, top + row * cell
            draw.rectangle((x, y, x + cell, y + cell), fill=cell_color(value), outline=(225, 225, 225))
            if abs(value) >= 0.45 or row == col:
                text_color = "white" if abs(value) >= 0.58 else "black"
                draw.text((x + 5, y + 16), f"{value:.2f}", font=small, fill=text_color)
    for col, label in enumerate(labels):
        x = left + col * cell + 5
        draw.text((x, top - 32), label[:9], font=small, fill="black")
    draw.text((left, height - 42), "Blue: positive correlation    Red: negative correlation    Values shown when |r| >= 0.45", font=regular, fill="black")
    image.save(HEATMAP)


def make_report_pages(
    target_corr: pd.Series,
    target_spearman: pd.Series,
    pearson: pd.DataFrame,
) -> list[Path]:
    """Rasterize Korean report pages so Korean glyphs render consistently in Word previews."""
    width, height = 1700, 2200
    font_path = "/System/Library/Fonts/Supplemental/AppleGothic.ttf"
    title_font = ImageFont.truetype(font_path, 54)
    heading_font = ImageFont.truetype(font_path, 32)
    body_font = ImageFont.truetype(font_path, 23)
    small_font = ImageFont.truetype(font_path, 19)
    bold_font = ImageFont.truetype(font_path, 25)

    def canvas() -> tuple[Image.Image, ImageDraw.ImageDraw]:
        page = Image.new("RGB", (width, height), "white")
        return page, ImageDraw.Draw(page)

    def wrapped(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, max_chars: int = 48, line_gap: int = 12, font=body_font) -> int:
        lines: list[str] = []
        current = ""
        for word in text.split(" "):
            candidate = f"{current} {word}".strip()
            if len(candidate) > max_chars:
                if current:
                    lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        for line in lines:
            draw.text((x, y), line, font=font, fill="black")
            y += font.size + line_gap
        return y

    def header(draw: ImageDraw.ImageDraw, text: str) -> int:
        draw.text((110, 100), text, font=title_font, fill="black")
        draw.line((110, 180, 1590, 180), fill=(31, 78, 121), width=4)
        return 245

    def grid(draw: ImageDraw.ImageDraw, x: int, y: int, headers: list[str], rows: list[list[str]], widths: list[int]) -> int:
        row_h = 60
        for row_index, row in enumerate([headers] + rows):
            xx = x
            fill = (31, 78, 121) if row_index == 0 else ((243, 247, 250) if row_index % 2 == 0 else (255, 255, 255))
            color = "white" if row_index == 0 else "black"
            for value, cell_width in zip(row, widths):
                draw.rectangle((xx, y, xx + cell_width, y + row_h), fill=fill, outline=(210, 210, 210), width=1)
                draw.text((xx + 12, y + 16), value, font=small_font, fill=color)
                xx += cell_width
            y += row_h
        return y

    pages: list[Path] = []
    page, draw = canvas()
    y = header(draw, "ZeroFest 변수간 상관분석 보고서")
    draw.text((110, y), "다음 30분 판매량 예측을 위한 합성 축제 운영 데이터 분석", font=body_font, fill=(80, 80, 80))
    y += 90
    draw.text((110, y), "분석 결론", font=heading_font, fill="black")
    y += 60
    y = wrapped(draw, "다음 30분 판매량은 최근 30분 판매량과 가장 강한 양의 상관을 보였다. 이전 30분 판매량도 높은 양의 상관을 보였다. 강수확률과 가격은 음의 상관을, 할인율과 캠퍼스 규모는 양의 상관을 보였다.", 110, y)
    y += 30
    draw.text((110, y), "분석 기준", font=heading_font, fill="black")
    y += 60
    y = grid(draw, 110, y, ["항목", "내용"], [
        ["분석 데이터", "2023-2025년 3개 가상 대학 3일 축제 38,880건"],
        ["관측 해상도", "1분 관측 단위와 초 단위 ISO 8601 타임스탬프"],
        ["종속변수", "future_sales_30m 다음 30분 판매량"],
        ["분석 방법", "Pearson 선형 상관계수와 Spearman 순위 상관계수"],
        ["해석 범위", "모델 개발용 합성 데이터 분석 결과"],
    ], [360, 1120])
    y += 70
    draw.text((110, y), "핵심 시사점", font=heading_font, fill="black")
    y += 60
    for item in [
        "판매 추세 변수는 수요 예측의 핵심 입력으로 유지한다.",
        "강수확률과 할인율은 실시간 운영 Action과 연결해 관리한다.",
        "시간 변수 중 일부는 같은 정보를 중복 표현하므로 피처 축소 검증이 필요하다.",
    ]:
        draw.text((125, y), "• " + item, font=body_font, fill="black")
        y += 48
    path = OUTPUT_DIR / "correlation_report_page_1.png"
    page.save(path)
    pages.append(path)

    page, draw = canvas()
    y = header(draw, "변수간 상관 히트맵")
    draw.text((110, y), "Pearson 상관계수 기준", font=body_font, fill=(80, 80, 80))
    chart = Image.open(HEATMAP).convert("RGB")
    chart.thumbnail((1450, 1320))
    page.paste(chart, ((width - chart.width) // 2, y + 45))
    y += chart.height + 105
    y = wrapped(draw, "파란색은 양의 상관, 빨간색은 음의 상관을 뜻한다. 진한 색일수록 변수 사이의 선형 관계가 강하다. 대각선은 같은 변수의 자기상관이다.", 110, y, max_chars=57)
    y += 40
    draw.text((110, y), "중복성 확인", font=heading_font, fill="black")
    y += 60
    y = wrapped(draw, "분 단위 시각과 당일 종료까지 남은 시간은 같은 운영 시간을 다른 방식으로 나타내므로 절대 상관이 매우 높다. 축제 일차와 전체 축제 종료까지 남은 시간도 높은 음의 상관을 보인다.", 110, y, max_chars=57)
    path = OUTPUT_DIR / "correlation_report_page_2.png"
    page.save(path)
    pages.append(path)

    page, draw = canvas()
    y = header(draw, "종속변수와의 관계")
    draw.text((110, y), "종속변수 future_sales_30m 다음 30분 판매량", font=body_font, fill=(80, 80, 80))
    y += 75
    rows = []
    for feature, value in target_corr.head(10).items():
        relation = "양의 관계" if value > 0 else "음의 관계"
        rows.append([LABELS[feature], f"{value:+.3f}", f"{target_spearman[feature]:+.3f}", relation])
    y = grid(draw, 110, y, ["변수", "Pearson", "Spearman", "관계"], rows, [540, 260, 280, 420])
    y += 65
    draw.text((110, y), "모델 관리 제안", font=heading_font, fill="black")
    y += 55
    for item in [
        "최근 판매량과 이전 판매량은 유지한다.",
        "minute_of_day, minutes_to_close, seconds_to_close는 대표 변수 선택 실험을 수행한다.",
        "피처 제거 전후의 시간 순서 Hold out MAE와 R²를 비교해 최종 조합을 결정한다.",
        "second_of_minute는 예측 신호보다 판매 이벤트 감사와 외부 데이터 정렬에 사용한다.",
    ]:
        y = wrapped(draw, "• " + item, 125, y, max_chars=58)
        y += 8
    path = OUTPUT_DIR / "correlation_report_page_3.png"
    page.save(path)
    pages.append(path)
    return pages


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    history = pd.read_csv(HISTORY_PATH)
    history["category_code"] = history["menu_category"].map({"food": 0, "drink": 1, "dessert": 2})
    history["price_k"] = history["price"] / 1000
    history["campus_scale"] = history["student_count"] / 10000
    columns = FEATURES + ["future_sales_30m"]
    numeric = history[columns].apply(pd.to_numeric, errors="coerce")
    pearson = numeric.corr(method="pearson")
    spearman = numeric.corr(method="spearman")
    target = "future_sales_30m"
    target_corr = pearson[target].drop(target).sort_values(key=lambda series: series.abs(), ascending=False)
    target_spearman = spearman[target].drop(target)
    build_heatmap(pearson)

    pages = make_report_pages(target_corr, target_spearman, pearson)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.12)
    section.bottom_margin = Inches(0.12)
    section.left_margin = Inches(0.15)
    section.right_margin = Inches(0.15)
    for index, page in enumerate(pages):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        if index:
            paragraph.paragraph_format.page_break_before = True
        paragraph.add_run().add_picture(str(page), width=Inches(7.75))
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Created {HEATMAP}")
    return

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "AppleGothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "AppleGothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "AppleGothic")
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("ZeroFest 변수간 상관분석 보고서")
    subtitle = document.add_paragraph("다음 30분 판매량 예측을 위한 합성 축제 운영 데이터 분석")
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    add_heading(document, "분석 결론", 1)
    add_body(
        document,
        "다음 30분 판매량은 최근 30분 판매량과 가장 강한 양의 상관을 보였다. 이전 30분 판매량도 높은 양의 상관을 보였다. 강수확률과 가격은 음의 상관을, 할인율과 캠퍼스 규모는 양의 상관을 보였다.",
        "핵심 결과. ",
    )
    add_body(
        document,
        "시간 변수 사이에는 구조적으로 높은 중복성이 있다. 특히 분 단위 시각과 종료까지 남은 시간은 같은 운영 시간을 다른 방식으로 표현한다. 학습 성능은 트리 모델이 유지할 수 있으나, 설명 가능성과 데이터 관리 측면에서는 중복 피처를 통제해야 한다.",
        "의사결정 시사점. ",
    )

    add_heading(document, "데이터와 방법", 1)
    add_table(document, ["항목", "내용"], [
        ["분석 데이터", f"2023부터 2025년까지 3개 가상 대학의 3일 축제 이력 {len(history):,}건"],
        ["관측 해상도", "1분 관측 단위와 초 단위 ISO 8601 타임스탬프"],
        ["종속변수", "future_sales_30m 다음 30분 판매량"],
        ["독립변수", f"판매 추세 시간 날씨 공연 할인 메뉴 가격 캠퍼스 규모 등 {len(FEATURES)}개"],
        ["분석 방법", "Pearson 선형 상관계수와 Spearman 순위 상관계수 비교"],
        ["해석 범위", "모든 값은 합성 데이터 기반의 모델 개발용 결과이며 실제 판매 실증값이 아님"],
    ])

    document.add_page_break()
    add_heading(document, "변수간 상관 히트맵", 1)
    add_body(document, "색이 진할수록 절대 상관이 크다. 파란색은 양의 상관, 빨간색은 음의 상관을 뜻한다. 숫자는 절대값 0.45 이상인 관계와 대각선 자기상관에 표시했다.")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(HEATMAP), width=Inches(6.75))
    caption = document.add_paragraph("그림 1 Pearson 상관계수 히트맵")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)

    document.add_page_break()
    add_heading(document, "종속변수와의 관계", 1)
    target_rows = []
    for feature, value in target_corr.head(10).items():
        direction = "양의 관계" if value > 0 else "음의 관계"
        strength = "강함" if abs(value) >= 0.7 else "중간" if abs(value) >= 0.3 else "약함"
        target_rows.append([
            LABELS[feature], f"{value:+.3f}", f"{target_spearman[feature]:+.3f}", f"{strength} {direction}",
        ])
    add_table(document, ["변수", "Pearson", "Spearman", "해석"], target_rows)

    add_heading(document, "중복 변수 점검", 1)
    pairs: list[tuple[str, str, float]] = []
    for left_index, left in enumerate(FEATURES):
        for right in FEATURES[left_index + 1:]:
            pairs.append((left, right, pearson.loc[left, right]))
    pair_rows = []
    for left, right, value in sorted(pairs, key=lambda item: abs(item[2]), reverse=True)[:6]:
        action = "대표 변수 1개 선택 권장" if abs(value) >= 0.95 else "모델 중요도와 성능 확인 권장"
        pair_rows.append([LABELS[left], LABELS[right], f"{value:+.3f}", action])
    add_table(document, ["변수 1", "변수 2", "Pearson", "관리 제안"], pair_rows)

    add_heading(document, "운영 적용 제안", 1)
    add_body(document, "최근 30분 판매량과 이전 30분 판매량은 예측 입력으로 유지한다. 강수확률과 할인율은 운영 Action에 직접 연결할 수 있으므로 실시간 갱신 주기와 수집 품질을 관리한다.")
    add_body(document, "분 단위 시각, 종료까지 남은 분, 종료까지 남은 초는 같은 시간 정보를 겹쳐 담고 있다. 다음 모델 개선 시에는 대표 변수 선택, 피처 중요도 비교, Hold out 성능 차이 검증을 통해 최종 조합을 결정한다.")
    add_body(document, "초 단위 타임스탬프는 판매 이벤트 감사와 공연 및 날씨 데이터 정렬에 사용한다. 초 자체의 상관은 낮으므로 단독 수요 신호로 해석하지 않는다.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("ZeroFest AI  |  Sample Synthetic Festival Dataset")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Created {HEATMAP}")


if __name__ == "__main__":
    main()
