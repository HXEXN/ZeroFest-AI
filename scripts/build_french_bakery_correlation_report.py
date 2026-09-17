"""Build a correlation heatmap report from Kaggle French Bakery Daily Sales."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATASET_HANDLE = "matthieugimbert/french-bakery-daily-sales"
DATA_PATH = Path.home() / ".cache/kagglehub/datasets/matthieugimbert/french-bakery-daily-sales/versions/1/Bakery sales.csv"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT = OUTPUT_DIR / "French_Bakery_Daily_Sales_상관분석_히트맵_보고서.docx"
HEATMAP = OUTPUT_DIR / "french_bakery_daily_correlation_heatmap.png"

LABELS = {
    "sales_quantity": "일 판매수량",
    "revenue_eur": "일 매출 EUR",
    "ticket_count": "고객 티켓 수",
    "unique_articles": "판매 품목 수",
    "avg_unit_price": "평균 단가 EUR",
    "avg_items_ticket": "티켓당 수량",
    "avg_sale_hour": "평균 판매시각",
    "weekend": "주말 여부",
    "weekday": "요일",
    "month": "월",
    "day_index": "경과 일수",
}
HEATMAP_COLUMNS = list(LABELS)


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype("/System/Library/Fonts/Supplemental/AppleGothic.ttf", size)


def wrap(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, max_chars: int = 51, size: int = 25) -> int:
    words, line, lines = text.split(" "), "", []
    for word in words:
        candidate = f"{line} {word}".strip()
        if len(candidate) > max_chars and line:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    for item in lines:
        draw.text((x, y), item, font=font(size), fill="black")
        y += size + 14
    return y


def table(draw: ImageDraw.ImageDraw, y: int, headers: list[str], rows: list[list[str]], widths: list[int]) -> int:
    x, row_h = 110, 66
    for index, row in enumerate([headers] + rows):
        xx = x
        fill = (28, 70, 112) if index == 0 else ((244, 248, 252) if index % 2 == 0 else "white")
        color = "white" if index == 0 else "black"
        for value, width in zip(row, widths):
            draw.rectangle((xx, y, xx + width, y + row_h), fill=fill, outline=(208, 216, 224))
            draw.text((xx + 12, y + 18), value, font=font(18), fill=color)
            xx += width
        y += row_h
    return y


def page(title: str) -> tuple[Image.Image, ImageDraw.ImageDraw, int]:
    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    draw.text((110, 100), title, font=font(50), fill="black")
    return image, draw, 215


def build_daily_frame() -> tuple[pd.DataFrame, pd.Timestamp, pd.Timestamp, int]:
    source = pd.read_csv(DATA_PATH)
    source["date"] = pd.to_datetime(source["date"])
    source["hour"] = source["time"].str.split(":").str[0].astype(int)
    source["price_eur"] = source["unit_price"].str.replace(" €", "", regex=False).str.replace(",", ".", regex=False).astype(float)
    source["revenue_eur"] = source["Quantity"] * source["price_eur"]

    daily = source.groupby("date").agg(
        sales_quantity=("Quantity", "sum"),
        revenue_eur=("revenue_eur", "sum"),
        ticket_count=("ticket_number", "nunique"),
        unique_articles=("article", "nunique"),
        avg_unit_price=("price_eur", "mean"),
        avg_items_ticket=("Quantity", "sum"),
        avg_sale_hour=("hour", "mean"),
    )
    daily["avg_items_ticket"] = daily["avg_items_ticket"] / daily["ticket_count"]
    daily["weekend"] = (daily.index.weekday >= 5).astype(int)
    daily["weekday"] = daily.index.weekday
    daily["month"] = daily.index.month
    daily["day_index"] = range(len(daily))
    return daily, source["date"].min(), source["date"].max(), len(source)


def ensure_data_path() -> Path:
    """Download the Kaggle dataset once when it is absent from the local cache."""
    if DATA_PATH.exists():
        return DATA_PATH
    import kagglehub

    dataset_dir = Path(kagglehub.dataset_download(DATASET_HANDLE))
    candidates = list(dataset_dir.glob("*.csv"))
    if len(candidates) != 1:
        raise FileNotFoundError("Expected exactly one CSV file in the Kaggle dataset directory")
    return candidates[0]


def build_heatmap(corr: pd.DataFrame) -> None:
    labels, cell, left, top = list(corr.columns), 101, 445, 150
    width, height = left + cell * len(labels) + 80, top + cell * len(labels) + 155
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((left, 38), "French Bakery Daily Sales Pearson Correlation Heatmap", font=font(27), fill="black")

    def color(value: float) -> tuple[int, int, int]:
        intensity = abs(value)
        base = (42, 102, 168) if value >= 0 else (190, 57, 67)
        return tuple(int(255 * (1 - intensity) + channel * intensity) for channel in base)

    for row, name in enumerate(labels):
        draw.text((15, top + row * cell + 34), LABELS[name], font=font(19), fill="black")
        for col in range(len(labels)):
            value = float(corr.iat[row, col])
            x, y = left + col * cell, top + row * cell
            draw.rectangle((x, y, x + cell, y + cell), fill=color(value), outline=(220, 220, 220))
            draw.text((x + 22, y + 39), f"{value:.2f}", font=font(16), fill="white" if abs(value) >= 0.54 else "black")
    for col, name in enumerate(labels):
        draw.text((left + col * cell + 7, top - 35), LABELS[name][:7], font=font(14), fill="black")
    draw.text((left, height - 49), "파란색은 양의 상관, 빨간색은 음의 상관", font=font(19), fill="black")
    image.save(HEATMAP)


def report_pages(daily: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, line_count: int) -> list[Path]:
    corr = daily[HEATMAP_COLUMNS].corr(method="pearson")
    spearman = daily[HEATMAP_COLUMNS].corr(method="spearman")
    build_heatmap(corr)
    target = corr["sales_quantity"].drop("sales_quantity").sort_values(key=lambda value: value.abs(), ascending=False)
    pages: list[Path] = []

    image, draw, y = page("French Bakery Daily Sales 변수간 상관분석 보고서")
    draw.text((110, y), "Kaggle Matthieu Gimbert French Bakery Daily Sales 기반", font=font(25), fill=(80, 80, 80))
    y += 90
    draw.text((110, y), "분석 결론", font=font(32), fill="black")
    y += 60
    y = wrap(draw, f"총 {line_count:,}개 판매 라인을 일 단위로 집계한 결과, 일 판매수량은 고객 티켓 수와 {target['ticket_count']:+.3f}, 판매 품목 수와 {target['unique_articles']:+.3f}의 강한 양의 상관을 보였다. 평균 판매시각은 {target['avg_sale_hour']:+.3f}로 음의 관계를 보였다.", 110, y)
    y += 30
    draw.text((110, y), "분석 범위", font=font(32), fill="black")
    y += 60
    y = table(draw, y, ["항목", "내용"], [
        ["원천 데이터", "matthieugimbert/french-bakery-daily-sales"],
        ["관측 기간", f"{start.date()} ~ {end.date()}"],
        ["분석 단위", f"일별 집계 {len(daily):,}일"],
        ["종속변수", "sales_quantity 일 판매수량"],
        ["독립변수", "고객 티켓 수 품목 수 평균 단가 티켓당 수량 시간 요일 월"],
        ["분석 방법", "Pearson 선형 상관계수 및 Spearman 순위 상관계수"],
    ], [365, 1110])
    y += 70
    draw.text((110, y), "해석 원칙", font=font(32), fill="black")
    y += 60
    wrap(draw, "상관계수는 동시 변화의 정도이며 인과관계를 뜻하지 않는다. 일 매출과 고객 티켓 수는 당일 판매 결과에서 계산되는 사후 지표이므로 행사 전 수요 예측모델의 입력값으로 그대로 사용하면 정보 누수가 발생할 수 있다.", 110, y)
    p = OUTPUT_DIR / "french_bakery_report_page_1.png"
    image.save(p); pages.append(p)

    image, draw, y = page("일별 운영 지표 상관 히트맵")
    draw.text((110, y), "Pearson 상관계수 기준", font=font(25), fill=(80, 80, 80))
    chart = Image.open(HEATMAP).convert("RGB")
    chart.thumbnail((1470, 1470))
    image.paste(chart, ((1700 - chart.width) // 2, y + 35))
    y += chart.height + 100
    wrap(draw, "일 판매수량과 고객 티켓 수, 판매 품목 수, 티켓당 수량은 함께 증가하는 패턴을 보인다. 평균 판매시각이 늦어질수록 일 판매수량이 낮은 경향은 조기 마감 또는 판매가 이른 시간대에 집중되는 운영 패턴과 관련될 수 있다.", 110, y)
    p = OUTPUT_DIR / "french_bakery_report_page_2.png"
    image.save(p); pages.append(p)

    image, draw, y = page("일 판매수량과의 관계 및 ZeroFest 적용")
    draw.text((110, y), "종속변수 sales_quantity 일 판매수량", font=font(25), fill=(80, 80, 80))
    y += 72
    rows = []
    for key, value in target.items():
        rows.append([LABELS[key], f"{value:+.3f}", f"{spearman.loc[key, 'sales_quantity']:+.3f}", "양의 관계" if value > 0 else "음의 관계"])
    y = table(draw, y, ["변수", "Pearson", "Spearman", "방향"], rows, [520, 300, 300, 400])
    y += 70
    draw.text((110, y), "모델 설계 시사점", font=font(32), fill="black")
    y += 60
    for statement in [
        "행사 전 예측에는 요일, 월, 사전 편성 공연, 날씨, 부스 메뉴와 가격처럼 사전에 아는 변수만 사용한다.",
        "행사 중 재예측에는 누적 티켓 수와 누적 판매량을 시점별 피처로 추가하되, 예측 기준 시각 이후 정보는 사용하지 않는다.",
        "일 매출과 당일 최종 티켓 수는 목표값과 동시 산출되는 결과이므로 학습 입력에서 분리한다.",
        "축제 데이터는 판매와 재고 변화를 분 단위 타임스탬프로 기록해 품절 시점과 할인 효과를 별도 검증한다.",
    ]:
        y = wrap(draw, statement, 110, y, max_chars=57)
        y += 8
    p = OUTPUT_DIR / "french_bakery_report_page_3.png"
    image.save(p); pages.append(p)
    return pages


def package_docx(pages: list[Path]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.12)
    section.bottom_margin = Inches(0.12)
    section.left_margin = Inches(0.15)
    section.right_margin = Inches(0.15)
    for index, path in enumerate(pages):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if index:
            paragraph.paragraph_format.page_break_before = True
        paragraph.add_run().add_picture(str(path), width=Inches(7.75))
    document.save(OUTPUT)


def main() -> None:
    global DATA_PATH
    DATA_PATH = ensure_data_path()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    daily, start, end, lines = build_daily_frame()
    pages = report_pages(daily, start, end, lines)
    package_docx(pages)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
